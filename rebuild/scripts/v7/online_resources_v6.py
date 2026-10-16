#!/usr/bin/env python3
"""Online Resources for manuscript V6 (JIIM).

Writes
  manuscript_v6/Online_Resource_1.docx   Tables S1 to S9 and Figure S1
  manuscript_v6/Online_Resource_2.xlsx   per-method means and per-pair agreement
  manuscript_v6/figures/FigS1_matrices.{png,tif,pdf}

Every number is read from a results JSON (rebuild/results/v7/*.json,
results/v7/audit/audit_v7.json, results/v7/rescore_local, results/kaggle*,
results/perimage32, results/exclusions_v5.json) or computed here from the
per-image score files through scripts/perimage_v7.py (and pools_v5 for the
U and F synthetic means, as in bench_v7). Computed values are asserted against
the stored JSON values wherever a stored counterpart exists.

Run:  nice -n 10 python3 scripts/v7/online_resources_v6.py
"""
import itertools
import json
import math
import os
import re
import sys

import numpy as np
from scipy.stats import rankdata

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.dirname(HERE)
ROOT = os.path.dirname(SCRIPTS)                       # .../rebuild
PROJ = os.path.dirname(ROOT)                          # .../wound-recovery
sys.path.insert(0, SCRIPTS)
sys.path.insert(0, ROOT)
import perimage_v7 as P7                              # noqa: E402
import pools_v5 as PV5                                # noqa: E402
from woundbench.methods import METHOD_FAMILY          # noqa: E402
from woundbench.deep import DEEP_SPECS                # noqa: E402

RES = os.path.join(ROOT, "results")
V7 = os.path.join(RES, "v7")
MS = os.path.join(PROJ, "manuscript_v6")
FIGDIR = os.path.join(MS, "figures")
os.makedirs(FIGDIR, exist_ok=True)


def jl(*p):
    return json.load(open(os.path.join(*p)))


ST = jl(V7, "stats_v7.json")
BE = jl(V7, "bench_v7.json")
PR = jl(V7, "practical_v7.json")
RO = jl(V7, "robust_v7.json")
AU = jl(V7, "audit", "audit_v7.json")
RPP = os.path.join(V7, "robust_pooled_v7.json")
RP = jl(RPP) if os.path.exists(RPP) else None
import excl_cfg                                                      # noqa: E402
EXCL = excl_cfg.EXCL

REFS = list(P7.REFS)
CONTROL = P7.CONTROL
DS = REFS + [CONTROL]
WOUND = list(P7.WOUND)
COLON = list(P7.COLON)
SYNTH = list(P7.SYNTH)
NICE = {"medetec": "Medetec", "co2wounds": "CO2Wounds-V2", "fuseg": "FUSeg", "azh": "AZH",
        "kvasir_disjoint": "Kvasir", "kvasir_sessile": "Kvasir-sessile", "cvc_clinicdb": "CVC-ClinicDB",
        "cvc_colondb": "CVC-ColonDB", "etis_larib": "ETIS-Larib", "ph2": "PH2", "tn3k_test": "TN3K",
        "azh_train_slice": "Control (AZH training slice)", "fuseg_r256": "FUSeg at 256 px"}
SHORTNICE = dict(NICE, azh_train_slice="Control")
MINUS = "−"

# ====================================================================== formatting


def _neg(s):
    if s.startswith("-"):
        if float(s) == 0:
            return s[1:]
        return MINUS + s[1:]
    return s


def f(x, d=3):
    """Fixed decimals, minus sign U+2212, no negative zero."""
    if x is None or (isinstance(x, float) and not math.isfinite(x)):
        return "n/a"
    return _neg(f"{x:.{d}f}")


def fs(x, d=3):
    """Signed (explicit plus) fixed decimals."""
    if x is None or not math.isfinite(x):
        return "n/a"
    s = f(x, d)
    return s if s.startswith(MINUS) or float(f"{x:.{d}f}") == 0 else "+" + s


def iv(a, b=None, d=3):
    if b is None:
        a, b = a
    return f"[{f(a, d)}, {f(b, d)}]"


def eiv(e, d=3):
    """dict with est and (lo, hi) or ci."""
    if "ci" in e:
        return f"{f(e['est'], d)} {iv(e['ci'], d=d)}"
    return f"{f(e['est'], d)} {iv(e['lo'], e['hi'], d)}"


def fp(p, min_p=None):
    """p value with 2 significant digits; '(min)' when p is the smallest attainable."""
    if p is None:
        return "n/a"
    s = f"{p:.2g}"
    if "e" in s:
        m, e = s.split("e")
        s = f"{p:.{-int(e) + 1}f}"
    # keep trailing zero for 2 significant digits (e.g. 0.00040)
    digits = re.sub(r"^0\.0*", "", s.replace("-", ""))
    if len(digits) < 2:
        s = s + "0"
    if min_p is not None and abs(p - min_p) < 1e-12:
        s += " (min)"
    return s


def pc(x, d=1):
    return "n/a" if x is None else _neg(f"{100 * x:.{d}f}") + "%"


def pciv(v):
    return f"[{pc(v[0])}, {pc(v[1])}]"


assert fp(2 / 13860, 2 / 13860) == "0.00014 (min)"
assert fp(0.32886) == "0.33" and fp(0.0003968, 0.0003968) == "0.00040 (min)", fp(0.0003968, 0.0003968)
assert fp(0.005772) == "0.0058"
assert f(-0.0004) == "0.000" and f(-0.5) == MINUS + "0.500"

# ====================================================================== methods
G = {k: list(v) for k, v in P7.GROUPS.items()}
SEEDCOLS = P7.SEEDCOLS
ORIG10, NEW8 = P7.ORIG10, P7.NEW8
SETDEF = {k: v["groups"] for k, v in ST["definitions"]["method_sets"].items()}
SETN = {k: v["n"] for k, v in ST["definitions"]["method_sets"].items()}


def set_members(s):
    return sum((G[g] for g in SETDEF[s]), [])


for s in SETDEF:
    assert len(set_members(s)) == SETN[s], s
assert SETDEF["ALL"] == ["U", "F", "K", "M19"]
PAIRSETS = ["U", "F", "K", "M", "M11s", "M19", "N8s", "U+F", "U+K", "U+M", "U+M19", "ALL"]
assert all(s in SETDEF for s in PAIRSETS)

TR2 = jl(RES, "kaggle2", "train_report_k2.json")
AUROWS = {r["name"]: r for r in AU["check4_models"]["rows"]}
THR_K1P = jl(RES, "kaggle", "thresholds_pretrained.json")
THR_K1K = jl(RES, "kaggle", "thresholds_kvasir.json")
THR_K2S = jl(RES, "kaggle2", "thresholds_seeds.json")
THR_K2M = jl(RES, "kaggle2", "thresholds_modern2.json")
THR_DEEP = jl(RES, "deep_thresholds.json")
NNHELD = jl(RES, "kaggle2", "nnunet_heldout.json")
TR1 = jl(RES, "kaggle", "train_report_pretrained.json")
SPEC = {s[0]: (s[1], s[2]) for s in DEEP_SPECS}
ENC = {"resnet18": "ResNet-18", "resnet50": "ResNet-50", "mit_b0": "MiT-B0", "mit_b2": "MiT-B2",
       "mobilenet_v2": "MobileNetV2", "efficientnet-b0": "EfficientNet-B0",
       "tu-convnext_tiny.fb_in1k": "ConvNeXt-T", "tu-swin_tiny_patch4_window7_224": "Swin-T",
       "tu-efficientnet_b3": "EfficientNet-B3", "tu-regnety_016": "RegNetY-016"}
DEC = {"Unet": "U-Net", "UnetPlusPlus": "U-Net++", "DeepLabV3Plus": "DeepLabV3+", "FPN": "FPN",
       "PSPNet": "PSPNet", "Linknet": "LinkNet", "MAnet": "MA-Net", "Segformer": "SegFormer",
       "UPerNet": "UPerNet"}
LEARNED_DESC = {"logreg_all": "pixel classifier, logistic regression (all pixel features)",
                "gbm_all": "pixel classifier, gradient boosting (all pixel features)",
                "logreg_colour": "pixel classifier, logistic regression (colour features)",
                "unet": "small U-Net"}


def seed_name(base, seed, new):
    return P7._seed_name(base, seed, new)


def arch_of_modern(base, new):
    """decoder and encoder from the run-2 training report (seed 1 exists for all 18)."""
    r = TR2[seed_name(base, 1, new)]
    return DEC[r["arch"]], ENC[r["encoder"]]


# ====================================================================== per-method means (OR2)
METRICS = ("dice", "nsd")
MEANS = {m: {} for m in METRICS}          # metric -> set -> {name: mean}
NKEPT = {}
for d in DS:
    for m in METRICS:
        D = P7.load(d, m)
        MEANS[m][d] = dict(zip(D["names"], D["X"].mean(0)))
        NKEPT[d] = len(D["ids"])
        assert NKEPT[d] == ST["n_images"][d], d

# U and F against pools_v5 (independent route, V5 file reader)
TT = PV5.load()
for m in METRICS:
    for d in DS:
        for n in G["U"] + G["F"] + G["K"] + G["M"]:
            assert abs(MEANS[m][d][n] - TT[m][d][n]) < 1e-12, (m, d, n)
# group mean Dice against the stats_v7 floor table
for g in ("U", "F", "K", "M", "S", "M11s", "M19", "N8s"):
    for d in DS:
        v = np.mean([MEANS["dice"][d][n] for n in G[g]])
        assert abs(v - ST["floor"]["mean_dice"][g][d]) < 1e-12, (g, d)
# per-method seed-averaged means against the audit's independent recomputation
for d in ("medetec", "azh"):
    for n, r in AU["check8_recompute"][d]["means"].items():
        assert abs(MEANS["dice"][d][n] - r["audit"]) < 1e-12, (d, n)

# fuseg_r256: rescore files (P0), fuseg exclusions by index; U and F from rescore_local
R256 = "fuseg_r256"
keep = np.ones(120, bool)
for e in EXCL["fuseg"]:
    keep[e[0]] = False
NKEPT[R256] = int(keep.sum())
for m in METRICS:
    MEANS[m][R256] = {}
mi_ = {"dice": 0, "nsd": 2}
loc = jl(V7, "rescore_local", "fuseg_r256.json")
assert loc["variants"][0] == "P0" and loc["metrics"][0] == "dice" and loc["metrics"][2] == "nsd2"
assert loc["ids"] == jl(RES, "perimage32", "fuseg.json")["ids"] and len(loc["ids"]) == 120
Vl = np.asarray(loc["var"], float)
for k, n in enumerate(loc["names"]):
    for m in METRICS:
        MEANS[m][R256][n] = float(Vl[keep, k, 0, mi_[m]].mean())
indiv256 = {}
for fam in ("pretrained_s0", "pretrained_s1", "pretrained_s2", "modern2_s0", "modern2_s1", "modern2_s2",
            "kvasir_nets", "kvasir_unet", "nnunet"):
    j = jl(RES, "kaggle2", "rescore", fam, R256 + ".json")
    assert j["set"] == R256 and j["variants"][0] == "P0" and j["metrics"][2] == "nsd2" and j["n"] == 120
    Vv = np.asarray(j["var"], float)
    for k, n in enumerate(j["names"]):
        indiv256[n] = Vv[keep, k, 0, :]
for c, mem in SEEDCOLS.items():
    indiv256[c] = np.mean(np.stack([indiv256[x] for x in mem], 0), 0)
for n, A in indiv256.items():
    for m in METRICS:
        MEANS[m][R256][n] = float(A[:, mi_[m]].mean())
for fam in ("sam_v2_oracle", "sam_v2_nnbox"):
    j = jl(RES, "kaggle2", "scores", fam, R256 + ".json")
    assert j["n"] == 120
    for m in METRICS:
        A = np.asarray(j[m], float)
        for k, n in enumerate(j["names"]):
            MEANS[m][R256][n] = float(A[keep, k].mean())
# check against robust_v7 resolution block (M19 and K11 group means)
K11 = RO["definitions"]["groups"]["K11"]["members"]
for g, mem in (("M19", G["M19"]), ("K11", K11)):
    for m, rk in (("dice", "dice"), ("nsd", "nsd2")):
        v = np.mean([MEANS[m][R256][n] for n in mem])
        assert abs(v - RO["resolution"][f"{g}|{rk}"]["fuseg_r256_mean"]) < 1e-9, (g, m, v)

# synthetic sets: trained Kaggle models from load_synth, SAM run-1 from its score files,
# U and F Dice from pools_v5 (no per-image synthetic scores exist for them, as in bench_v7)
for s in SYNTH:
    for m in METRICS:
        Dd = P7.load_synth(s, m)
        MEANS[m][s] = dict(zip(Dd["names"], Dd["X"].mean(0)))
        j = jl(RES, "kaggle", "scores", "sam", s + ".json")
        A = np.asarray(j[m], float)
        for k, n in enumerate(j["names"]):
            MEANS[m][s][n] = float(A[:, k].mean())
    for n in G["U"] + G["F"]:
        MEANS["dice"][s][n] = float(TT["dice"][s][n])
    NKEPT[s] = len(P7.load_synth(s)["X"])
for g in ("U", "F", "K", "M", "S", "M11s", "M19", "N8s"):
    for s in SYNTH:
        v = np.mean([MEANS["dice"][s][n] for n in G[g]])
        assert abs(v - BE["mean_dice"][g][s]) < 1e-12, (g, s, v, BE["mean_dice"][g][s])

ALLSETS = DS + [R256] + SYNTH

# ====================================================================== per-pair agreement


def spearman(a, b):
    ra, rb = rankdata(a), rankdata(b)
    ra -= ra.mean()
    rb -= rb.mean()
    return float((ra * rb).sum() / math.sqrt((ra * ra).sum() * (rb * rb).sum()))


PAIRS = list(itertools.combinations(DS, 2))
assert len(PAIRS) == 66
AGREE = {}                                  # (metric, set, a, b) -> rho
n_checked = 0
for m in METRICS:
    for s in PAIRSETS:
        mem = set_members(s)
        stored = ST["disattenuated"][f"{m}|{s}"]["pairs"]
        for a, b in PAIRS:
            r = spearman([MEANS[m][a][n] for n in mem], [MEANS[m][b][n] for n in mem])
            AGREE[(m, s, a, b)] = r
            assert abs(r - stored[f"{a}~{b}"]["observed_full"]) < 1e-12, (m, s, a, b)
            n_checked += 1
assert n_checked == 2 * len(PAIRSETS) * 66
# summaries against stats_v7 point estimates
for m in METRICS:
    for s in PAIRSETS:
        wp = np.mean([AGREE[(m, s, a, b)] for a, b in itertools.combinations(WOUND, 2)])
        assert abs(wp - ST["point"][f"{m}|{s}"]["wound pairs"]) < 1e-12


def pair_category(a, b):
    if CONTROL in (a, b):
        return "control" if {a, b} == {"azh", CONTROL} else "with control"
    if P7.MODALITY[a] != P7.MODALITY[b]:
        return "across modality"
    if P7.SOURCE[a] == P7.SOURCE[b]:
        return "same modality, same source"
    return "same modality, different source"


print(f"means and pair agreement ready; {n_checked} pair values equal stats_v7.json", flush=True)

# ====================================================================== Figure S1
import matplotlib                                     # noqa: E402
matplotlib.use("Agg")
import matplotlib.pyplot as plt                       # noqa: E402

MM = 1 / 25.4
INK, SEC = "#111111", "#444444"
plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 8, "axes.edgecolor": SEC,
                     "axes.labelcolor": INK, "xtick.color": SEC, "ytick.color": SEC,
                     "axes.linewidth": 0.6, "xtick.major.width": 0.6, "ytick.major.width": 0.6,
                     "pdf.fonttype": 42})
FIGS1_SETS = [("U", "a  U: 18 untrained"), ("M19", "b  M19: 19 modern"), ("U+M19", "c  U+M19: 37 pooled")]
for s, t in FIGS1_SETS:
    assert t.split(": ")[1].split()[0] == str(SETN[s]), (s, t)


def matrix(s, m="dice"):
    M = np.full((len(REFS), len(REFS)), np.nan)
    for i, a in enumerate(REFS):
        for j, b in enumerate(REFS):
            if i != j:
                M[i, j] = AGREE[(m, s, a, b) if (m, s, a, b) in AGREE else (m, s, b, a)]
    return M


FIGS1_M = {s: matrix(s) for s, _ in FIGS1_SETS}
# explicit spot check of three cells against stats_v7.json (pair values stored there)
for s, a, b in (("U", "medetec", "fuseg"), ("M19", "co2wounds", "azh"), ("U+M19", "fuseg", "kvasir_disjoint")):
    i, j = REFS.index(a), REFS.index(b)
    assert FIGS1_M[s][i, j] == FIGS1_M[s][j, i]
    assert abs(FIGS1_M[s][i, j] - ST["disattenuated"][f"dice|{s}"]["pairs"][f"{a}~{b}"]["observed_full"]) < 1e-12
# and the wound-pair mean of each matrix against the stored point estimate
for s, _ in FIGS1_SETS:
    w = [FIGS1_M[s][REFS.index(a), REFS.index(b)] for a, b in itertools.combinations(WOUND, 2)]
    assert abs(np.mean(w) - ST["point"][f"dice|{s}"]["wound pairs"]) < 1e-12


def cell_text(v):
    """two decimals, leading zero dropped so that cells do not collide (0.54 -> .54)."""
    t = f(v, 2)
    return t.replace("0.", ".", 1) if t.lstrip(MINUS).startswith("0.") else t


def fig_s1():
    lab = [NICE[r] for r in REFS]
    fig, axes = plt.subplots(1, 3, figsize=(174 * MM, 72 * MM))
    for ax, (s, t) in zip(axes, FIGS1_SETS):
        M = FIGS1_M[s]
        shade = np.where(np.isnan(M), 1.0, np.abs(M))
        ax.imshow(shade, cmap="Greys", vmin=0, vmax=1.6)
        for i, j in itertools.product(range(len(REFS)), range(len(REFS))):
            if i != j:
                v = M[i, j]
                ax.text(j, i, cell_text(v), ha="center", va="center", fontsize=3.9,
                        color="white" if abs(v) > 0.8 else INK,
                        fontweight="bold" if float(f"{v:.2f}") < 0 else "normal")
        ax.set_xticks(range(len(REFS)))
        ax.set_yticks(range(len(REFS)))
        ax.set_xticklabels(lab, rotation=60, ha="right", fontsize=5.8)
        ax.set_yticklabels(lab if ax is axes[0] else [], fontsize=5.8)
        for b in (3.5, 8.5, 9.5):
            ax.axhline(b, color=INK, lw=0.8)
            ax.axvline(b, color=INK, lw=0.8)
        ax.set_title(t, fontsize=8, loc="left", color=INK)
    fig.tight_layout(w_pad=0.5)
    base = os.path.join(FIGDIR, "FigS1_matrices")
    fig.savefig(base + ".tif", dpi=600, pil_kwargs={"compression": "tiff_lzw"}, facecolor="white")
    fig.savefig(base + ".png", dpi=300, facecolor="white")
    fig.savefig(base + ".pdf", facecolor="white")
    plt.close(fig)
    return base + ".png"


FIGS1_PNG = fig_s1()
print("Figure S1 written", flush=True)

# ====================================================================== method rows (Table S1 and OR2)
GROUPLAB = {"U": "U", "F": "F", "K": "K", "P10": "Modern, original 10", "N8": "Modern, new 8",
            "NN": "nnU-Net", "S": "SAM/MedSAM"}


def modern_rows():
    rows = []
    for new, bases in ((False, ORIG10), (True, NEW8)):
        for b in bases:
            dec, enc = arch_of_modern(b, new)
            for sd in (0, 1, 2):
                n = seed_name(b, sd, new)
                if n in MEANS["dice"]["medetec"]:
                    rows.append({"name": n, "group": "N8" if new else "P10", "kind": "individual model",
                                 "arch": f"{dec}, {enc}", "seed": str(sd)})
    for new, bases in ((False, ORIG10), (True, NEW8)):
        for b in bases:
            dec, enc = arch_of_modern(b, new)
            c = f"{b}_in@seeds" if not new else f"{b}@seeds"
            sds = [str(int(re.search(r"_s(\d)$", x).group(1))) if re.search(r"_s\d$", x) else "0" for x in SEEDCOLS[c]]
            rows.append({"name": c, "group": "N8" if new else "P10", "kind": "seed average",
                         "arch": f"{dec}, {enc}", "seed": "mean of " + ", ".join(sds)})
    return rows


def all_rows():
    rows = [{"name": n, "group": "U", "kind": "method", "arch": f"classical, {METHOD_FAMILY[n]}", "seed": ""}
            for n in G["U"]]
    for g, suf in (("F", ""), ("K", "_kv")):
        for n in G[g]:
            b = n[:-len(suf)] if suf else n
            if b in SPEC:
                dec, enc = SPEC[b]
                a = f"{DEC[dec]}, {ENC[enc]} (random initialisation)"
            else:
                a = LEARNED_DESC[b]
            rows.append({"name": n, "group": g, "kind": "method", "arch": a, "seed": ""})
    rows += modern_rows()
    rows.append({"name": "nnunet_2d", "group": "NN", "kind": "individual model", "arch": "nnU-Net, 2D configuration",
                 "seed": "one training"})
    for n, a in (("sam_vitb_box", "SAM ViT-B, reference-box prompt (oracle)"),
                 ("medsam_vitb_box", "MedSAM ViT-B, reference-box prompt (oracle)"),
                 ("medsam_fixed_gtbox", "MedSAM ViT-B, official preprocessing, reference-box prompt (oracle)"),
                 ("sam_nnbox", "SAM ViT-B, prompted with nnU-Net boxes"),
                 ("medsam_fixed_nnbox", "MedSAM ViT-B, official preprocessing, prompted with nnU-Net boxes")):
        rows.append({"name": n, "group": "S", "kind": "box-prompted model", "arch": a, "seed": ""})
    return rows


MROWS = all_rows()
assert sorted(r["name"] for r in MROWS) == sorted(MEANS["dice"]["medetec"]), \
    set(MEANS["dice"]["medetec"]) ^ {r["name"] for r in MROWS}
assert len({r["name"] for r in MROWS}) == len(MROWS)

# ====================================================================== Online Resource 2
from openpyxl import Workbook                          # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side   # noqa: E402
from openpyxl.utils import get_column_letter          # noqa: E402

XF = "Times New Roman"
HFILL = PatternFill("solid", fgColor="E6E6E6")
HFONT = Font(name=XF, bold=True, size=10, color="000000")
BFONT = Font(name=XF, size=10, color="000000")
TFONT = Font(name=XF, bold=True, size=13, color="000000")
THIN = Border(bottom=Side(style="thin", color="8C8C8C"))


def set_kind(s):
    if s in REFS:
        return "clinical reference"
    if s == CONTROL:
        return "within-dataset control"
    if s == R256:
        return "resolution check (FUSeg test images, long side 256 px)"
    return "synthetic"


def synth_desc(s):
    m = re.match(r"genB_amp(\d+)_s(\d+)", s)
    if m:
        return f"generator B, texture amplitude {m.group(1)}, synthetic seed {m.group(2)}"
    m = re.match(r"(nr\w+)_s(\d+)", s)
    gen = {"nrA": "generator A, uniform drape", "nrB0": "generator B, texture ablated",
           "nrC0": "generator C, no noise", "nrC34": "generator C, noise 3.4"}[m.group(1)]
    return f"{gen}, synthetic seed {m.group(2)}"


def write_or2(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    lines = [
        ("Online Resource 2", TFONT),
        ("Method pool, clinical reference and training seed shape the measured validity of synthetic segmentation "
         "benchmarks: a study in wound imaging", Font(name=XF, italic=True, size=11)),
        ("", BFONT),
        ("Sheets", HFONT),
        ("Image sets: every image set, its kind, modality, and the number of images (and resampling units) the means use.", BFONT),
        ("Mean Dice, Mean NSD: per-method mean score on every image set. One row per method; every individual trained "
         "model (each training seed) and every seed average has its own row.", BFONT),
        ("Pair agreement: Spearman agreement between two image sets over a method set, for every pair of the 11 clinical "
         "references and the control (66 pairs), Dice and NSD.", BFONT),
        ("", BFONT),
        ("Definitions", HFONT),
        ("Mean score: mean over the kept images of the per-image score. Dice of an empty reference mask is 1 if the "
         "prediction is also empty and 0 otherwise. NSD is the normalised surface Dice at a tolerance of 2 px.", BFONT),
        ("Kept images: the clinical sets exclude 8 FUSeg test images that duplicate FUSeg training images and 47 control "
         "images that are crops of FUSeg photographs (Online Resource 1, Table S9). FUSeg at 256 px uses the FUSeg "
         "exclusions. Synthetic sets have no exclusions.".replace("8 FUSeg", f"{len(EXCL['fuseg'])} FUSeg")
         .replace("47 control", f"{len(EXCL['azh_train_slice'])} control"), BFONT),
        ("Resampling units (used by the bootstrap intervals of Online Resource 1): photographs for Medetec, AZH and the "
         "control (several stored images can come from one photograph), single images otherwise.", BFONT),
        ("Seed average (rows marked 'seed average'): per image, the mean of the per-image scores of the training seeds "
         "that exist, then the mean over images; it equals the mean of the seed models' means.", BFONT),
        ("Agreement: Spearman correlation (Pearson correlation of average ranks) between the per-method mean scores on "
         "dataset A and on dataset B, over the methods of the method set.", BFONT),
        ("Method sets: " + "; ".join(f"{s} ({SETN[s]} methods: {' + '.join(SETDEF[s])})" for s in PAIRSETS) +
         ". U: untrained classical methods. F: trained from scratch on FUSeg. K: the same 14 trained on Kvasir-SEG. "
         "M: the 10 original ImageNet-pretrained architectures at training seed 0 plus nnU-Net. M11s: the same 10 "
         "architectures as seed averages plus nnU-Net. M19: the 18 ImageNet-pretrained architectures as seed averages "
         "plus nnU-Net. N8s: the 8 new architectures as seed averages. SAM and MedSAM rows are never part of a method set.", BFONT),
        ("Pair category: same modality, same source (FUSeg and AZH; Kvasir and Kvasir-sessile; CVC-ClinicDB and "
         "CVC-ColonDB); same modality, different source; across modality; control (AZH with its training slice); with "
         "control (any other dataset with the control).", BFONT),
        ("", BFONT),
        ("Empty cells (no score exists)", HFONT),
        ("Synthetic sets, NSD of U and F: these methods were scored on synthetic images for Dice only; no per-image "
         "synthetic NSD exists for them.", BFONT),
        ("Synthetic sets, SAM/MedSAM with official preprocessing and the nnU-Net-box variants: not run on synthetic images.", BFONT),
        ("FUSeg at 256 px: the three Kvasir-trained pixel classifiers (logreg_all_kv, gbm_all_kv, logreg_colour_kv) and "
         "the first-run SAM and MedSAM were not run on it.", BFONT),
        ("", BFONT),
        ("Sources", HFONT),
        ("Clinical sets and control: per-image scores of rebuild/results/perimage32 (U, F) and rebuild/results/kaggle*/scores "
         "(trained models, SAM/MedSAM) through rebuild/scripts/perimage_v7.py. nnU-Net: first-session scores.", BFONT),
        ("FUSeg at 256 px: rescore files (pipeline P0) in rebuild/results/kaggle2/rescore (trained models; nnU-Net there is the "
         "second-session re-prediction) and rebuild/results/v7/rescore_local (U, F); SAM/MedSAM from rebuild/results/kaggle2/scores.", BFONT),
        ("Synthetic sets: per-image scores of the trained models and of the first-run SAM/MedSAM; U and F Dice means from "
         "rebuild/scripts/pools_v5.py, as in the benchmark analysis.", BFONT),
        ("Checks run by rebuild/scripts/v7/online_resources_v6.py before writing: every pair value equals the value stored "
         "in rebuild/results/v7/stats_v7.json (to 1e-12); U, F, K and M clinical means equal pools_v5; group mean Dice equals "
         "the stored floor and benchmark tables; FUSeg-at-256-px group means equal rebuild/results/v7/robust_v7.json.", BFONT),
        ("Values are unrounded numbers; the displayed format shows 3 decimals.", BFONT),
    ]
    for i, (t, ft) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = ft
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 150

    # image sets
    ws = wb.create_sheet("Image sets")
    hdr = ["set key", "name", "kind", "modality or generator", "images used", "resampling units"]
    ws.append(hdr)
    for s in ALLSETS:
        if s in DS:
            mod = P7.MODALITY.get(s, "wound")
            units = ST["n_units"][s]
            name = NICE[s]
        elif s == R256:
            mod, units, name = "wound", NKEPT[s], NICE[s]
        else:
            mod, units, name = synth_desc(s), NKEPT[s], s
        ws.append([s, name, set_kind(s), mod, NKEPT[s], units])
    style_sheet(ws, len(hdr), widths=[22, 30, 46, 52, 13, 16], numfmt={5: "0", 6: "0"})

    # means
    for m, title in (("dice", "Mean Dice"), ("nsd", "Mean NSD")):
        ws = wb.create_sheet(title)
        hdr = ["method", "group", "model type", "architecture or algorithm", "training seed"] + ALLSETS
        ws.append(hdr)
        for r in MROWS:
            vals = [MEANS[m][s].get(r["name"]) for s in ALLSETS]
            ws.append([r["name"], GROUPLAB[r["group"]], r["kind"], r["arch"], r["seed"]] +
                      [None if v is None else float(v) for v in vals])
        style_sheet(ws, len(hdr), widths=[26, 20, 18, 58, 16] + [14] * len(ALLSETS),
                    numfmt={k: "0.000" for k in range(6, len(hdr) + 1)})
        ws.freeze_panes = "F2"

    ws = wb.create_sheet("Pair agreement")
    hdr = ["metric", "method set", "methods in set", "dataset A", "dataset B", "pair category", "Spearman agreement"]
    ws.append(hdr)
    for m in METRICS:
        for s in PAIRSETS:
            for a, b in PAIRS:
                ws.append(["Dice" if m == "dice" else "NSD", s, SETN[s], NICE[a], NICE[b], pair_category(a, b),
                           AGREE[(m, s, a, b)]])
    style_sheet(ws, len(hdr), widths=[8, 12, 15, 30, 30, 32, 20], numfmt={3: "0", 7: "0.000"})
    ws.freeze_panes = "A2"
    wb.save(path)


def style_sheet(ws, ncol, widths, numfmt):
    for j in range(1, ncol + 1):
        c = ws.cell(row=1, column=j)
        c.font, c.fill, c.border = HFONT, HFILL, THIN
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BFONT
            if c.column in numfmt and isinstance(c.value, (int, float)):
                c.number_format = numfmt[c.column]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


OR2 = os.path.join(MS, "Online_Resource_2.xlsx")
write_or2(OR2)
print("Online Resource 2 written", flush=True)

# ====================================================================== docx engine
from docx import Document                                       # noqa: E402
from docx.enum.section import WD_ORIENT                         # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT                  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK         # noqa: E402
from docx.oxml import OxmlElement                               # noqa: E402
from docx.oxml.ns import qn                                     # noqa: E402
from docx.shared import Cm, Pt, RGBColor                        # noqa: E402

FONT = "Times New Roman"
INKC = RGBColor(0, 0, 0)
GREYC = RGBColor(0x3F, 0x3F, 0x3F)
LINKC = "28598A"
RULE, INNER, HEAD = "8C8C8C", "CFCFCF", "E6E6E6"
TEXT_LOG = []                     # every string written, for the final dash check


class Doc:
    def __init__(self):
        self.d = Document()
        st = self.d.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(10)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        sec = self.d.sections[0]
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
        for k in ("left_margin", "right_margin"):
            setattr(sec, k, Cm(1.8))
        sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(2.0)
        sec.footer_distance = Cm(0.8)
        self.width_cm = 29.7 - 3.6
        self.bm_id = 0
        self.toc = []
        self._footer_page_numbers(sec)

    def _footer_page_numbers(self, sec):
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        for t, txt in (("begin", None), (None, "PAGE"), ("end", None)):
            if t:
                e = OxmlElement("w:fldChar")
                e.set(qn("w:fldCharType"), t)
            else:
                e = OxmlElement("w:instrText")
                e.set(qn("xml:space"), "preserve")
                e.text = txt
            r._r.append(e)
        r.font.size = Pt(9)
        r.font.name = FONT

    def run(self, p, text, size=10, bold=False, italic=False, color=INKC):
        TEXT_LOG.append(text)
        r = p.add_run(text)
        r.font.name, r.font.size, r.bold, r.italic = FONT, Pt(size), bold, italic
        r.font.color.rgb = color
        return r

    def para(self, text="", size=10, bold=False, italic=False, color=INKC, after=4, before=0,
             align=WD_ALIGN_PARAGRAPH.LEFT, keep=False):
        p = self.d.add_paragraph()
        p.alignment = align
        pf = p.paragraph_format
        pf.space_after, pf.space_before = Pt(after), Pt(before)
        pf.keep_with_next = keep
        if text:
            self.run(p, text, size, bold, italic, color)
        return p

    def bookmark(self, p, name):
        self.bm_id += 1
        s = OxmlElement("w:bookmarkStart")
        s.set(qn("w:id"), str(self.bm_id))
        s.set(qn("w:name"), name)
        e = OxmlElement("w:bookmarkEnd")
        e.set(qn("w:id"), str(self.bm_id))
        p._p.insert(0, s)
        p._p.append(e)

    def link(self, p, text, anchor, size=10):
        TEXT_LOG.append(text)
        h = OxmlElement("w:hyperlink")
        h.set(qn("w:anchor"), anchor)
        h.set(qn("w:history"), "1")
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        f_ = OxmlElement("w:rFonts")
        for a in ("w:ascii", "w:hAnsi", "w:cs"):
            f_.set(qn(a), FONT)
        c = OxmlElement("w:color")
        c.set(qn("w:val"), LINKC)
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size * 2)))
        for x in (f_, c, u, sz):
            rpr.append(x)
        r.append(rpr)
        t = OxmlElement("w:t")
        t.text = text
        t.set(qn("xml:space"), "preserve")
        r.append(t)
        h.append(r)
        p._p.append(h)

    def page_break(self):
        p = self.d.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)

    def caption(self, label, title, text, bm=None, before=10):
        bm = label.replace(" ", "")
        self.toc.append((label, title, bm))
        p = self.para(before=before, after=4, keep=True)
        p.paragraph_format.line_spacing = 1.05
        self.run(p, f"{label}. ", 9.5, bold=True)
        self.run(p, title + ". ", 9.5, bold=True)
        self.run(p, text, 9.5, color=GREYC)
        if bm:
            self.bookmark(p, bm)
        return p

    def note(self, text):
        p = self.para(text, size=8.5, color=GREYC, after=6, before=2)
        return p

    def table(self, header, rows, widths, size=8, left_cols=1, bold_rows=()):
        tot = sum(widths)
        widths = [w * min(1.0, self.width_cm / tot) for w in widths]
        t = self.d.add_table(rows=1 + len(rows), cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        tblPr = t._tbl.tblPr
        b = OxmlElement("w:tblBorders")
        for edge, val, col, sz in (("top", "single", RULE, 8), ("bottom", "single", RULE, 8),
                                   ("left", "nil", None, 0), ("right", "nil", None, 0),
                                   ("insideH", "single", INNER, 2), ("insideV", "nil", None, 0)):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), val)
            if col:
                e.set(qn("w:sz"), str(sz))
                e.set(qn("w:color"), col)
                e.set(qn("w:space"), "0")
            b.append(e)
        tblPr.append(b)
        mar = OxmlElement("w:tblCellMar")
        for side, v in (("top", 25), ("bottom", 25), ("left", 55), ("right", 55)):
            e = OxmlElement(f"w:{side}")
            e.set(qn("w:w"), str(v))
            e.set(qn("w:type"), "dxa")
            mar.append(e)
        tblPr.append(mar)
        grid = t._tbl.tblGrid
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(w / 2.54 * 1440)))
        for ri, vals in enumerate([header] + rows):
            row = t.rows[ri]
            trPr = row._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit")
            trPr.append(cs)
            if ri == 0:
                th = OxmlElement("w:tblHeader")
                trPr.append(th)
            for ci, v in enumerate(vals):
                cell = row.cells[ci]
                cell.width = Cm(widths[ci])
                if ri == 0:
                    tcPr = cell._tc.get_or_add_tcPr()
                    sh = OxmlElement("w:shd")
                    sh.set(qn("w:val"), "clear")
                    sh.set(qn("w:color"), "auto")
                    sh.set(qn("w:fill"), HEAD)
                    tcPr.append(sh)
                lines = str(v).split("\n")
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci < left_cols else WD_ALIGN_PARAGRAPH.RIGHT
                if ri == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci < left_cols else WD_ALIGN_PARAGRAPH.RIGHT
                for li, line in enumerate(lines):
                    r = self.run(p, line, size, bold=(ri == 0 or (ri - 1) in bold_rows))
                    if li < len(lines) - 1:
                        r.add_break()
        return t

    def figure(self, png, width_cm):
        p = self.para(align=WD_ALIGN_PARAGRAPH.CENTER, after=2, before=6, keep=True)
        p.add_run().add_picture(png, width=Cm(width_cm))
        return p

    def save(self, path):
        self.d.save(path)


def m_h(b, mk="method_only_repro_B2000", hk="hierarchical"):
    return f"{f(b['est'])}\nm {iv(b[mk])}\nh {iv(b[hk])}"


def short(n):
    return n.replace("@seeds", "").replace("_in", "")

# ====================================================================== shared caption text
from math import comb                                            # noqa: E402

NW, NC = len(WOUND), len(COLON)
N_SAME = sum(1 for a, b in itertools.combinations(REFS, 2) if P7.MODALITY[a] == P7.MODALITY[b])
N_ACROSS = comb(len(REFS), 2) - N_SAME
N_SRC = sum(1 for a, b in itertools.combinations(REFS, 2) if P7.SOURCE[a] == P7.SOURCE[b])
SRC_PAIRS = [(a, b) for a, b in itertools.combinations(REFS, 2) if P7.SOURCE[a] == P7.SOURCE[b]]
assert (comb(NW, 2), comb(NC, 2), N_SAME, N_ACROSS, N_SRC) == (6, 10, 16, 39, 3)
MKEY = "method_only_repro_B2000"
B_M = int(re.search(r"B(\d+)$", MKEY).group(1))
B_H = ST["definitions"]["B_hier"]
N_RELAB = ST["exact"]["dice|U"]["full_55"]["n_relabelings"]
MIN_P = ST["exact"]["dice|U"]["full_55"]["min_p"]
assert abs(MIN_P - 2 / N_RELAB) < 1e-15
N_UNITS_TXT = ", ".join(f"{SHORTNICE[d]} {ST['n_units'][d]} units for {ST['n_images'][d]} images" for d in ("medetec", "azh", CONTROL))

SETS14 = ["U", "F", "K", "M", "M11s", "M19", "N8s", "U+F", "U+K", "U+M", "U+M11s", "U+M19", "ALL57", "ALL"]
assert sorted(SETS14) == sorted(SETDEF)
SETTXT = (f"Method sets: U, {SETN['U']} untrained classical methods; F, {SETN['F']} methods trained from scratch on FUSeg; "
          f"K, the same {SETN['K']} methods trained on Kvasir-SEG; M, the {len(ORIG10)} original ImageNet-pretrained "
          f"architectures at training seed 0 plus nnU-Net ({SETN['M']}); M11s, the same {len(ORIG10)} architectures as seed "
          f"averages plus nnU-Net ({SETN['M11s']}); M19, all {len(ORIG10) + len(NEW8)} ImageNet-pretrained architectures "
          f"({len(ORIG10)} original, {len(NEW8)} new) as seed averages plus nnU-Net ({SETN['M19']}); N8s, the {len(NEW8)} new "
          f"architectures as seed averages ({SETN['N8s']}); pooled sets U+F ({SETN['U+F']}), U+K ({SETN['U+K']}), "
          f"U+M ({SETN['U+M']}), U+M11s ({SETN['U+M11s']}), U+M19 ({SETN['U+M19']}); ALL57 = U+F+K+M ({SETN['ALL57']}); "
          f"ALL = U+F+K+M19 ({SETN['ALL']}). A seed average is one column per architecture: per image, the mean of the "
          f"scores of the training seeds that exist.")
AGRTXT = ("Agreement is the Spearman correlation of the per-method mean scores of two image sets over the methods of "
          "a method set.")
PAIRTXT = (f"Wound pairs: mean over the {comb(NW, 2)} pairs of the {NW} wound datasets. Colonoscopy pairs: mean over the "
           f"{comb(NC, 2)} pairs of the {NC} colonoscopy datasets. Modality difference ({comb(len(REFS), 2)} pairs): mean "
           f"agreement of the {N_SAME} same-modality pairs minus the mean of the {N_ACROSS} across-modality pairs of the "
           f"{len(REFS)} clinical references; {comb(len(REFS), 2) - N_SRC}-pair version: the same without the {N_SRC} "
           f"same-source pairs ({'; '.join(f'{NICE[a]} with {NICE[b]}' for a, b in SRC_PAIRS)}). Control: AZH with its "
           "training slice.")
UNITTXT = (f"Resampling units are photographs for Medetec, AZH and the control ({N_UNITS_TXT}) and "
           "single images otherwise.")
MHTXT = (f"m: method bootstrap (methods resampled with replacement within each group, pooled sets keep their "
         f"composition), B = {B_M:,}; h: hierarchical bootstrap (methods within group and resampling units within "
         f"dataset), B = {B_H:,}; both 95% percentile intervals.")
EXTXT = (f"Exact p: share of the {N_RELAB:,} distinct relabelings of the {len(REFS)} modality labels whose modality "
         f"difference is at least the observed one; (min) marks the smallest attainable p (2/{N_RELAB:,} = "
         f"{fp(MIN_P)}).")
DSNAMES = ", ".join(NICE[r] for r in REFS)

# ====================================================================== S1


def table_s1(D):
    rows = []
    tr2_saved = [r for r in AU["check4_models"]["rows"] if r.get("source") == "run 2"]
    run2_all = [r for r in AU["check4_models"]["rows"] if r.get("source") == "run 2" or "diverged" in r["status"]]
    diverged = [r["name"] for r in AU["check4_models"]["rows"] if "diverged" in r["status"]]
    assert len(run2_all) == 44 and len(tr2_saved) == 41 and len(diverged) == 3, (len(run2_all), len(tr2_saved))
    for n in G["U"]:
        rows.append([n, "U", f"classical, {METHOD_FAMILY[n]}", "none", "n/a", "n/a", "n/a", "n/a", "n/a"])
    for g, suf, data, thr in (("F", "", "FUSeg training images, from scratch", THR_DEEP),
                              ("K", "_kv", "Kvasir-SEG training images, from scratch", THR_K1K)):
        for n in G[g]:
            b = n[:-len(suf)] if suf else n
            if b in SPEC:
                dec, enc = SPEC[b]
                a = f"{DEC[dec]}, {ENC[enc]}"
            else:
                a = LEARNED_DESC[b]
            th = f(thr[n], 2) if n in thr else "nr"
            rows.append([n, g, a, data, "one training", th, "nr", "n/a", "n/a"])
    ntr = {TR2[k]["n_train"] for k in TR2}
    nsw = {TR2[k]["n_sweep"] for k in TR2}
    ep = {TR2[k]["epochs"] for k in TR2}
    assert len(ntr) == len(nsw) == len(ep) == 1
    ntr, nsw, ep = ntr.pop(), nsw.pop(), ep.pop()
    for new, bases in ((False, ORIG10), (True, NEW8)):
        for b in bases:
            dec, enc = arch_of_modern(b, new)
            have, ths, hd = [], [], []
            for sd in (0, 1, 2):
                n = seed_name(b, sd, new)
                r = AUROWS[n]
                if "diverged" in r["status"]:
                    ths.append("div.")
                    hd.append("diverged")
                    continue
                have.append(str(sd))
                # thresholds and held-out Dice: the audit table must equal the raw files
                tf = THR_K1P if (not new and sd == 0) else (THR_K2S if not new else THR_K2M)
                assert tf[n] == r["threshold"], n
                src = TR1[n]["sweep_dice"] if (not new and sd == 0) else TR2[n]["sweep_dice"]
                assert src == r["heldout_dice"], n
                ths.append(f(r["threshold"], 2))
                hd.append(f(r["heldout_dice"]))
            miss = [str(sd) for sd in (0, 1, 2) if str(sd) not in have]
            seeds = ", ".join(have) + (f" (seed {', '.join(miss)} diverged)" if miss else "")
            key = f"{b}@seeds" if new else f"{b}_in"
            rows.append([key, "modern, new 8" if new else "modern, original 10", f"{dec}, {enc}",
                         f"FUSeg training split ({ntr} images), ImageNet initialisation", seeds, " / ".join(ths)] + hd)
    rows.append(["nnunet_2d", "nnU-Net", "nnU-Net, 2D configuration", "FUSeg training split", "one training",
                 f(NNHELD["threshold"], 2) + " (fixed)", f(NNHELD["mean_dice"]), "n/a", "n/a"])
    assert NNHELD["n"] == nsw
    for n, a in (("sam_vitb_box", "SAM ViT-B, reference box"), ("medsam_vitb_box", "MedSAM ViT-B, reference box"),
                 ("medsam_fixed_gtbox", "MedSAM ViT-B, official preprocessing, reference box"),
                 ("sam_nnbox", "SAM ViT-B, nnU-Net boxes"),
                 ("medsam_fixed_nnbox", "MedSAM ViT-B, official preprocessing, nnU-Net boxes")):
        rows.append([n, "SAM/MedSAM", a, "none (published weights)", "n/a", "n/a", "n/a", "n/a", "n/a"])
    hdr = ["Method", "Group", "Architecture or algorithm", "Training data", "Training seeds available",
           "Threshold (seed 0 / 1 / 2)", "Held-out Dice, seed 0", "Seed 1", "Seed 2"]
    cap = ("Every method in the study. Group: U, untrained classical methods (family in the third column); F, trained "
           "from scratch on FUSeg; K, the same methods trained from scratch on Kvasir-SEG; modern, original 10 and new 8: "
           "ImageNet-pretrained architectures trained on FUSeg at training seeds 0, 1 and 2; nnU-Net; SAM/MedSAM: "
           "box-prompted foundation models, used only as oracle references and never pooled. Architecture: decoder, "
           "encoder. Modern models: the training split is the FUSeg training images minus duplicates; "
           f"{ep} epochs; the decision threshold of each model was chosen on {nsw} held-out FUSeg training images, and "
           "'Held-out Dice' is the mean Dice on those images at that threshold (optimistic, because the threshold was "
           "tuned on the same images). nnU-Net: one training, fixed threshold, held-out Dice on the same "
           f"{nsw} images, which it never saw. Seed 0 of the original 10 was trained in a first GPU session, seeds 1 and 2 "
           f"and all seeds of the new 8 in a second session ({len(tr2_saved)} of {len(run2_all)} second-session trainings "
           f"saved; {', '.join(diverged)} diverged with non-finite values under mixed precision and were not imputed; "
           "Table S9). The modern rows name the seed-0 key of an original architecture or the seed-average key of a new "
           "one; individual models carry the suffix _s1 or _s2 (original 10) or _s0, _s1, _s2 (new 8) in Online "
           "Resource 2. n/a: not applicable. nr: not stored in a results file. div.: diverged.")
    D.caption("Table S1", "Methods, training and held-out Dice", cap, bm="S1", before=0)
    D.table(hdr, rows, [3.6, 2.4, 5.6, 4.6, 2.6, 2.6, 1.6, 1.4, 1.4], size=7.5, left_cols=5)
    return len(rows)


# ====================================================================== S2


def table_s2(D):
    counts = {}
    LV = [("same dataset", "Control"), ("same source", f"Same source ({N_SRC})"),
          ("different source, same modality", f"Different source, same modality ({N_SAME - N_SRC})"),
          ("across modality", f"Across modality ({N_ACROSS})"), ("same modality (16 pairs)", f"Same modality ({N_SAME})"),
          ("wound pairs", f"Wound pairs ({comb(NW, 2)})"), ("wound different-source pairs", "Wound, different source (5)"),
          ("fuseg~azh", "FUSeg with AZH"), ("colonoscopy pairs", f"Colonoscopy pairs ({comb(NC, 2)})"),
          ("modality_diff", "Modality difference (55)"), ("modality_diff_excl_within_source", "Modality difference (52)")]
    assert sum(1 for a, b in itertools.combinations(WOUND, 2) if P7.SOURCE[a] != P7.SOURCE[b]) == 5
    for part, m, mname in (("a", "dice", "Dice"), ("b", "nsd", "NSD")):
        rows = [[s, str(SETN[s])] + [f(ST["point"][f"{m}|{s}"][k]) for k, _ in LV] for s in SETS14]
        cap = (f"Point estimates on the full data. Each column is the mean agreement over the "
               "pairs of one category (number of pairs in brackets). Same source: two datasets from one source "
               f"({'; '.join(f'{NICE[a]} with {NICE[b]}' for a, b in SRC_PAIRS)}). {PAIRTXT} {AGRTXT} {SETTXT}")
        D.caption(f"Table S2{part}", f"Agreement levels by pair category ({mname})", cap,
                  bm="S2" if part == "a" else None, before=0 if part == "a" else 10)
        D.table(["Set", "n"] + [h for _, h in LV], rows, [1.5, 0.8] + [2.1] * len(LV), size=7.5, left_cols=1)
        counts[f"S2{part}"] = len(rows)
        D.page_break()
    Q = [("wound pairs", f"Wound pairs ({comb(NW, 2)})"),
         ("comp|wound|in x out", f"Wound with non-wound ({NW * (len(REFS) - NW)})"),
         ("colonoscopy pairs", f"Colonoscopy pairs ({comb(NC, 2)})"),
         ("modality_diff", "Modality difference (55)"), ("modality_diff_excl_within_source", "Modality difference (52)"),
         ("same dataset", "Control")]
    for part, m, mname in (("c", "dice", "Dice"), ("d", "nsd", "NSD")):
        rows = [[s, str(SETN[s])] + [m_h(ST["bootstrap"][f"{m}|{s}"][k]) for k, _ in Q] for s in SETS14]
        cap = (f"Each cell: estimate, then the m and h intervals. {MHTXT} "
               f"{UNITTXT} Resampling the units adds noise to every method mean and so pulls agreement down in each draw; "
               "h intervals therefore sit partly below the estimate. Wound with non-wound: mean over the "
               f"{NW * (len(REFS) - NW)} pairs of a wound dataset with a non-wound dataset. {PAIRTXT}")
        D.caption(f"Table S2{part}", f"Agreement with method and hierarchical bootstrap intervals ({mname})", cap)
        D.table(["Set", "n"] + [h for _, h in Q], rows, [1.5, 0.8] + [3.4] * len(Q), size=7.5)
        counts[f"S2{part}"] = len(rows)
        D.page_break()
    rows = []
    for m, mname in (("dice", "Dice"), ("nsd", "NSD")):
        for s in SETS14:
            e = ST["exact"][f"{m}|{s}"]
            a, b, lv = e["full_55"], e["excl_within_source_52"], e["levels_excl"]
            assert a["n_relabelings"] == b["n_relabelings"] == N_RELAB
            rows.append([mname, s, str(SETN[s]), f(a["diff"]), fp(a["p"], a["min_p"]), f(b["diff"]),
                         fp(b["p"], b["min_p"]), f(lv["different source, same modality (13)"]),
                         f(lv["across modality (39)"]), f(lv["same source (3, excluded)"])])
    cap = (f"Exact relabeling tests of the modality difference. {EXTXT} The 52-pair test drops the {N_SRC} same-source "
           "pairs under every relabeling. The last three columns give the category means that enter the 52-pair "
           f"difference and the mean of the dropped same-source pairs. {PAIRTXT} {SETTXT.split('. A seed')[0]}.")
    D.caption("Table S2e", "Exact tests of the modality difference (55 and 52 pairs)", cap)
    D.table(["Metric", "Set", "n", "Difference (55)", "Exact p (55)", "Difference (52)", "Exact p (52)",
             f"Different source, same modality ({N_SAME - N_SRC})", f"Across modality ({N_ACROSS})",
             f"Same source ({N_SRC}, dropped)"], rows, [1.4, 1.5, 0.8, 2.2, 2.4, 2.2, 2.4, 3.0, 2.4, 2.4], size=7.5,
            left_cols=2)
    counts["S2e"] = len(rows)
    D.page_break()
    rows = []
    for s in SETS14:
        v = ST["dice_noempty"][s]
        e = ST["exact"][f"dice_noempty|{s}"]
        assert abs(v["modality_diff"] - e["full_55"]["diff"]) < 1e-12 and abs(v["p"] - e["full_55"]["p"]) < 1e-12
        assert abs(v["excl_within_source"] - e["excl_within_source_52"]["diff"]) < 1e-12
        assert abs(v["dice_with_empty"] - ST["point"][f"dice|{s}"]["modality_diff"]) < 1e-12
        rows.append([s, str(SETN[s]), f(v["modality_diff"]), fp(e["full_55"]["p"], e["full_55"]["min_p"]),
                     f(v["excl_within_source"]), fp(e["excl_within_source_52"]["p"], e["excl_within_source_52"]["min_p"]),
                     f(v["dice_with_empty"]), f(v["wound pairs"]), f(v["wound pairs with empty"])])
    ne = ", ".join(f"{SHORTNICE[d]} {n}" for d, n in ST["n_empty"].items() if n)
    cap = ("Dice with images whose reference mask is empty excluded (per-method means recomputed without them; "
           f"empty reference masks: {ne}; all other datasets have none). With empties: the main-analysis value. "
           f"{EXTXT} {PAIRTXT}")
    D.caption("Table S2f", "Empty-mask-excluded Dice", cap)
    D.table(["Set", "n", "Modality difference (55), no empty", "Exact p", "Modality difference (52), no empty", "Exact p",
             "Modality difference (55), with empties", "Wound pairs, no empty", "Wound pairs, with empties"],
            rows, [1.6, 0.8, 2.9, 2.3, 2.9, 2.3, 2.9, 2.4, 2.4], size=7.5)
    counts["S2f"] = len(rows)
    return counts


# ====================================================================== S3


def table_s3(D):
    counts = {}
    LS = ["U", "M", "M11s", "M19", "N8s"]
    rows = []
    for m, mname in (("dice", "Dice"), ("nsd", "NSD")):
        for d in REFS:
            row = [mname, NICE[d]]
            for s in LS:
                e = ST["leave_one_dataset_out"][f"{m}|{s}"][d]["full"]
                assert e["n_pairs"] == comb(len(REFS) - 1, 2)
                row.append(f"{f(e['diff'])} ({fp(e['p'], e['min_p'])})")
            rows.append(row)
    nrel = sorted({ST["leave_one_dataset_out"][f"dice|M19"][d]["full"]["n_relabelings"] for d in REFS})
    cap = (f"Modality difference over the {comb(len(REFS) - 1, 2)} pairs of the remaining "
           f"{len(REFS) - 1} clinical references, with its exact relabeling p in brackets. The number of distinct "
           f"relabelings depends on the dropped dataset ({', '.join(f'{x:,}' for x in nrel)}); (min) marks the smallest "
           f"attainable p for that design. Sets: U ({SETN['U']}), M ({SETN['M']}), M11s ({SETN['M11s']}), "
           f"M19 ({SETN['M19']}), N8s ({SETN['N8s']}); definitions as in Table S2. {AGRTXT}")
    D.caption("Table S3a", "Leave one dataset out", cap, bm="S3", before=0)
    D.table(["Metric", "Dropped dataset"] + [f"{s}: difference (p)" for s in LS], rows,
            [1.4, 3.0] + [3.6] * len(LS), size=7.5, left_cols=2)
    counts["S3a"] = len(rows)
    D.page_break()
    rows = []
    full = {m: ST["point"][f"{m}|M19"] for m in METRICS}
    rows.append(["none (all 19)"] + [f(full[m][k]) for m in METRICS
                                     for k in ("modality_diff", "modality_diff_excl_within_source", "wound pairs")])
    for n in G["M19"]:
        rows.append([short(n)] + [f(ST["leave_one_method_out_M19"][m][n][k]) for m in METRICS
                                  for k in ("modality_diff", "modality_diff_excl_within_source", "wound pairs")])
    cap = (f"Agreement summaries over the remaining {SETN['M19'] - 1} methods. Method keys "
           "as in Table S1, seed averages shown by architecture key (for example unet_r18 is the seed average of the "
           f"three unet_r18 models). {PAIRTXT} {AGRTXT}")
    D.caption("Table S3b", "Leave one method out of M19", cap)
    D.table(["Dropped method", "Dice: modality difference (55)", "Dice: modality difference (52)", "Dice: wound pairs",
             "NSD: modality difference (55)", "NSD: modality difference (52)", "NSD: wound pairs"], rows,
            [3.6] + [3.0] * 6, size=7.5, bold_rows=(0,))
    counts["S3b"] = len(rows)
    return counts

# ====================================================================== S4
SDS = ST["seeds"]
SEEDSETS = ["P9", "N7", "C16"]
SB = ST["definitions"]["seed_analysis_boot"]
SEEDMH = (f"m: architecture bootstrap (architectures resampled with replacement; the same draw for all seeds and for "
          f"the seed average), B = {SB['B_method']:,}; h: hierarchical bootstrap (architectures and resampling units), "
          f"B = {SB['B_hier']:,}; both 95% percentile intervals.")


def seedsets_txt():
    s = SDS["sets"]
    for k in SEEDSETS:
        assert s[k]["n"] == len(s[k]["architectures"])
    excl = [c for c in G["M19"] if c not in s["C16"]["architectures"]]
    return (f"Architecture sets with all three training seeds: P9, the {s['P9']['n']} original architectures "
            f"({', '.join(short(c) for c in s['P9']['architectures'])}); N7, the {s['N7']['n']} new architectures "
            f"({', '.join(short(c) for c in s['N7']['architectures'])}); C16 = P9 + N7. Not in these sets: "
            f"{', '.join(short(c) for c in excl)} (missing seeds or one training). Seed 0 of the original architectures "
            "comes from the first GPU session and seeds 1 and 2 from the second, so for P9 seed and session effects "
            "cannot be separated.")


def table_s4(D):
    counts = {}
    for part, m, mname in (("a", "dice", "Dice"), ("b", "nsd", "NSD")):
        rows = []
        for d in DS:
            row = [SHORTNICE[d]]
            for k in SEEDSETS:
                p = SDS[f"{m}|{k}"]["per_dataset"][d]
                sb = p["spearman_brown_3"]
                row += [f(p["r_01"]), f(p["r_02"]), f(p["r_12"]), f(p["rbar"]), "undef." if sb is None else f(sb)]
            rows.append(row)
        hdr = ["Dataset"] + [f"{k} {c}" for k in SEEDSETS for c in ("r01", "r02", "r12", "r-bar", "SB3")]
        cap = (f"Same dataset, different training seed. r_st: Spearman agreement between the per-architecture "
               "mean scores of seed s and seed t on the same images. r-bar: mean of r01, r02 and r12. SB3: Spearman-Brown "
               "reliability of the 3-seed average, 3 r-bar / (1 + 2 r-bar), undefined (undef.) when r-bar is at most "
               f"−0.5; negative values mean no reliability. {seedsets_txt()}")
        D.caption(f"Table S4{part}", f"Ranking reliability across training seeds ({mname})", cap,
                  bm="S4" if part == "a" else None, before=0 if part == "a" else 10)
        D.table(hdr, rows, [3.0] + [1.42] * 15, size=7, left_cols=1)
        counts[f"S4{part}"] = len(rows)
    D.page_break()
    rows = []
    for m, mname in (("dice", "Dice"), ("nsd", "NSD")):
        for k in SEEDSETS:
            x = SDS[f"{m}|{k}"]["summary"]
            rows.append([mname, k, str(SDS["sets"][k]["n"])] +
                        [m_h(x[q], "method_boot", "hierarchical") for q in ("rbar_wound", "rbar_colon", "rbar_refs11",
                                                                             "rbar_control")])
    cap = (f"Mean r-bar (Table S4a, S4b) over dataset groups: the {NW} wound datasets, the {NC} colonoscopy datasets, all "
           f"{len(REFS)} clinical references, and the control. Each cell: estimate, then m and h intervals. {SEEDMH}")
    D.caption("Table S4c", "Seed-to-seed agreement on the same dataset, by dataset group", cap)
    D.table(["Metric", "Set", "n", f"Wound ({NW})", f"Colonoscopy ({NC})", f"All {len(REFS)} references", "Control"],
            rows, [1.4, 1.2, 0.8, 4.2, 4.2, 4.2, 4.2], size=7.5, left_cols=2)
    counts["S4c"] = len(rows)
    D.page_break()
    SQ = [("seed0|wound pairs", "Wound pairs, seed 0"), ("seed1|wound pairs", "Wound pairs, seed 1"),
          ("seed2|wound pairs", "Wound pairs, seed 2"), ("single|wound pairs", "Wound pairs, mean of single seeds"),
          ("avg|wound pairs", "Wound pairs, seed average"),
          ("avg_minus_single|wound pairs", "Seed average minus mean of single seeds"),
          ("rbar_wound", "Same wound dataset, different seed (r-bar)"),
          ("crossseed|wound pairs", "Different wound dataset, different seed"),
          ("gap|wound: same-dataset cross-seed minus cross-dataset cross-seed",
           "Gap: same dataset minus different dataset (both different seed)"),
          ("single|colonoscopy pairs", "Colonoscopy pairs, mean of single seeds"),
          ("avg|colonoscopy pairs", "Colonoscopy pairs, seed average"),
          ("single|modality_diff", "Modality difference, mean of single seeds"),
          ("avg|modality_diff", "Modality difference, seed average"),
          ("single|same dataset", "Control, mean of single seeds"), ("avg|same dataset", "Control, seed average")]
    for part, m, mname in (("d", "dice", "Dice"), ("e", "nsd", "NSD")):
        rows = [[lab] + [m_h(SDS[f"{m}|{k}"]["summary"][q], "method_boot", "hierarchical") for k in SEEDSETS]
                for q, lab in SQ]
        cap = (f"Cross-dataset agreement with single seeds and with seed averages. Seed s: every architecture "
               "represented by its seed-s model. Mean of single seeds: mean over seeds 0, 1 and 2. Seed average: the "
               "seed-averaged columns of the same architectures. Different wound dataset, different seed: seed s on one "
               f"wound dataset against seed t on another (s different from t; 6 ordered seed pairs times {comb(NW, 2)} "
               "wound pairs), the counterpart of the same-dataset r-bar, so that training noise enters both and only the "
               f"dataset differs. Gap: r-bar minus that value. Each cell: estimate, then m and h intervals. {SEEDMH} "
               f"{PAIRTXT}")
        D.caption(f"Table S4{part}", f"Single seeds, seed averages and the dataset gap ({mname})", cap)
        D.table(["Quantity"] + [f"{k} (n {SDS['sets'][k]['n']})" for k in SEEDSETS], rows, [7.0, 5.2, 5.2, 5.2],
                size=7.5)
        counts[f"S4{part}"] = len(rows)
        D.page_break()
    rows = []
    for k, v in SDS["named_groups"].items():
        m, g = k.split("|", 1)
        rows.append(["Dice" if m == "dice" else "NSD", g, str(v["n"]), f(v["wound pairs"]), f(v["colonoscopy pairs"]),
                     f(v["modality_diff"]), f(v["same dataset"])])
    cap = ("Point estimates. P10_sk: the original architectures at training seed k (models "
           "that exist); N8_sk: the new architectures at seed k; M (V6) = P10_s0 + nnU-Net; M_sk = P10_sk + nnU-Net; "
           "M19_sk = P10_sk + N8_sk + nnU-Net. nnU-Net has one training and appears in every single-seed set. Set sizes "
           f"differ because of the diverged trainings. Seed-averaged sets (M11s, M19, N8s) for comparison. {PAIRTXT} "
           f"{AGRTXT}")
    D.caption("Table S4f", "Single-seed method sets", cap)
    D.table(["Metric", "Set", "n", "Wound pairs", "Colonoscopy pairs", "Modality difference (55)", "Control"], rows,
            [1.4, 6.6, 0.8, 2.6, 2.6, 2.9, 2.4], size=7.5, left_cols=2)
    counts["S4f"] = len(rows)
    D.page_break()
    rows = []
    for m, mname in (("dice", "Dice"), ("nsd", "NSD")):
        for d, p in SDS["disattenuation"][f"{m}|C16"]["per_dataset"].items():
            rows.append([mname, SHORTNICE[d], f(p["R_img_seedavg"]), f(p["R_img_single_mean"]), f(p["R_seed_1"]),
                         f(p["R_seed_3"]), f(p["R_combined_avg_additive"]), f(p["R_combined_single_additive"]),
                         f(p["R_combined_avg_product"])])
    cap = ("Used to correct agreement for seed noise and image-sampling noise. Model: observed "
           "method mean = true value + seed error + image-sampling error. R_img: split-half reliability over "
           f"{ST['definitions']['n_splits']} random splits of the resampling units (Spearman-Brown median), for the "
           "seed-averaged columns and as the mean over seeds of the single-seed values. R_seed: seed reliability, r-bar "
           "for one seed and its Spearman-Brown value for 3 seeds. Additive rule: R = R_img + R_seed − 1; product "
           "rule: R = R_img R_seed. n/a: undefined. The additive rule over-corrects if part of the image error is "
           "specific to a seed; both rules are written for Pearson correlations and are approximate for Spearman.")
    D.caption("Table S4g", "Reliabilities for the noise correction (C16)", cap)
    D.table(["Metric", "Dataset", "R_img, seed average", "R_img, single seeds", "R_seed, 1 seed", "R_seed, 3 seeds",
             "R additive, seed average", "R additive, single seed", "R product, seed average"], rows,
            [1.4, 2.6] + [2.6] * 7, size=7.5, left_cols=2)
    counts["S4g"] = len(rows)
    D.page_break()
    rows = []
    for m, mname in (("dice", "Dice"), ("nsd", "NSD")):
        for pr, r in SDS["disattenuation"][f"{m}|C16"]["wound_pairs"].items():
            a, b = pr.split("~")
            rows.append([mname, f"{NICE[a]} with {NICE[b]}", f(r["observed_seedavg"]), f(r["observed_single_mean"]),
                         f(r["dis_img_only"]), f(r["dis_seed_only"]), f(r["dis_additive"]), f(r["dis_product"]),
                         f(r["dis_additive_single"])])
    cap = ("Observed: agreement of the seed averages and the mean over single "
           "seeds. Corrected: r / sqrt(R_A R_B) with R for image noise only, seed noise only, both (additive rule), both "
           "(product rule), and the additive rule applied to single seeds. Values above 1 occur when the reliabilities "
           "are small; they show that the correction is fragile, not that agreement exceeds 1.")
    D.caption("Table S4h", "Wound-pair agreement corrected for seed and image noise (C16)", cap)
    D.table(["Metric", "Pair", "Observed, seed average", "Observed, single seeds", "Image noise only",
             "Seed noise only", "Both, additive", "Both, product", "Additive, from single seeds"], rows,
            [1.4, 4.6] + [2.4] * 7, size=7.5, left_cols=2)
    counts["S4h"] = len(rows)
    rows = []
    for m, mname in (("dice", "Dice"), ("nsd", "NSD")):
        for k in SEEDSETS:
            for c, a in SDS["disattenuation"][f"{m}|{k}"]["categories"].items():
                def g_(key):
                    v = a[key]
                    return f"{f(v['raw_mean'])} / {f(v['capped_mean'])}\n({v['n_defined']} of {a['n_pairs']}; {v['n_abs_gt1']})"
                rows.append([mname, k, c.replace("azh~azh_train_slice", "AZH with control"),
                             f(a["observed_seedavg_mean"]), f(a["observed_single_mean"]), g_("dis_img_only"),
                             g_("dis_seed_only"), g_("dis_additive"), g_("dis_product"), g_("dis_additive_single")])
    cap = ("Corrected cells: raw mean / mean after capping each pair at "
           "−1 and 1; in brackets the number of pairs where the correction is defined (R above 0 for both datasets) "
           "of all pairs, and the number of pairs with absolute corrected value above 1. Columns as in Table S4h.")
    D.caption("Table S4i", "Corrected agreement by pair category", cap)
    D.table(["Metric", "Set", "Category", "Observed, seed average", "Observed, single seeds", "Image noise only",
             "Seed noise only", "Both, additive", "Both, product", "Additive, from single seeds"], rows,
            [1.3, 1.0, 3.6, 2.0, 2.0] + [2.6] * 5, size=7, left_cols=3)
    counts["S4i"] = len(rows)
    D.page_break()
    rows = []
    for m, mname in (("dice", "Dice"), ("nsd", "NSD")):
        for d in DS:
            p = SDS["spread"][f"{m}|{d}"]
            rows.append([mname, SHORTNICE[d], f(p["M19_mean"]), f(p["M19_range"]), f(p["M19_sd"]),
                         f(p["M19_median_adjacent_gap"]), f(p["C16_between_sd_of_seedavg"]), f(p["C16_seed_sd_pooled"]),
                         f(p["C16_seed_sd_median"]), f(p["C16_se_of_3seed_avg"]), f(p["ratio_between_sd_to_seed_sd"], 2),
                         f(p["C16_icc_single"]), f(p["C16_icc_3seed_avg"]), f(p["M19_image_se_median"])])
    cap = (f"Dice and NSD. M19 columns: mean, range and SD of the "
           f"{SETN['M19']} method means (seed averages) and the median gap between neighbours after sorting. Between SD: "
           "SD of the 16 seed-averaged C16 means. Seed SD: SD over the 3 seeds of one architecture's mean, pooled over "
           "C16 as the root mean variance, and its median; SE of a 3-seed average = pooled seed SD / sqrt(3). ICC: share "
           "of the between-architecture variance that is not seed noise, for one seed and for a 3-seed average "
           "(negative when the seed noise exceeds the observed spread). Image SE: median over M19 of the SD of a method "
           f"mean over {ST['definitions']['B_hier']:,} resamples of the units; it is largely shared by all methods, so "
           "it moves rankings less than its size suggests.")
    D.caption("Table S4j", "Spread of modern method means against seed noise", cap)
    D.table(["Metric", "Dataset", "M19 mean", "M19 range", "M19 SD", "Median neighbour gap", "Between SD (C16)",
             "Seed SD, pooled", "Seed SD, median", "SE of 3-seed average", "Between / seed SD", "ICC, one seed",
             "ICC, 3-seed average", "Image SE, median"], rows, [1.2, 2.0] + [1.75] * 12, size=7, left_cols=2)
    counts["S4j"] = len(rows)
    return counts

# ====================================================================== S5
BSETS = ["U", "F", "K", "M", "M11s", "M19", "P10s", "N8s", "U+F", "U+K", "U+M", "U+M11s", "U+M19", "U+F+K+M",
         "U+F+K+M19"]
BLAB = {"U+F+K+M": "ALL57", "U+F+K+M19": "ALL"}
for s in BSETS:
    if s in SETN:
        assert BE["pools"][s] == SETN[s], s
assert BE["pools"]["U+F+K+M"] == SETN["ALL57"] and BE["pools"]["U+F+K+M19"] == SETN["ALL"]
B_B = BE["B"]
BENCHTXT = ("Benchmark agreement (bench) with a wound reference R: Spearman agreement of mean Dice on generator B "
            "with R, averaged over the 3 synthetic seeds (2024, 7, 101); each synthetic set has "
            f"{BE['checks']['n_images_per_synth_set'][0]} images. Other: mean agreement of the other {NW - 1} wound datasets "
            "with R. Diff = bench minus other: positive means the benchmark ranks the methods more like R than another "
            f"wound dataset does. Mean: average over the {NW} wound references (other then equals the mean of the "
            f"{comb(NW, 2)} wound pairs). Dice only. ALL57 = U+F+K+M; ALL = U+F+K+M19; P10s: the {len(ORIG10)} original "
            "architectures as seed averages; other sets as in Table S2.")
BMB = (f"Intervals: method bootstrap (methods resampled with replacement within each group, pooled sets keep their "
       f"composition), B = {B_B:,}, 95% percentile.")
BHB = (f"Hierarchical bootstrap: methods within the group plus, independently per dataset, resampling units of the "
       f"clinical sets and images of the synthetic sets, B = {B_B:,}; data-only bootstrap: units and images only, "
       f"methods fixed, B = {B_B:,} (context, not the primary interval). Only K and the modern sets have per-image "
       "synthetic scores; U and F synthetic means exist per method only, so sets containing them get the method "
       "bootstrap only.")


def bl(s):
    return BLAB.get(s, s)


def table_s5(D):
    counts = {}
    rows = []
    for s in BSETS:
        a, b = BE["rel"][s]["40"], BE["rel"][s]["11"]
        assert a["mean|other"]["est"] == b["mean|other"]["est"]
        rows.append([bl(s), str(BE["pools"][s]), eiv(a["mean|bench"]), eiv(a["mean|diff"]), eiv(b["mean|bench"]),
                     eiv(b["mean|diff"]), eiv(a["mean|other"])])
    D.caption("Table S5a", "Benchmark validity by method set (mean over wound references)",
              f"Generator B at texture amplitude 40 and 11. {BENCHTXT} Each cell: estimate [interval]. {BMB}",
              bm="S5", before=0)
    D.table(["Set", "n", "Bench, amplitude 40", "Diff, amplitude 40", "Bench, amplitude 11", "Diff, amplitude 11",
             "Other (wound with wound)"], rows, [1.8, 0.8, 4.0, 4.0, 4.0, 4.0, 4.0], size=7.5)
    counts["S5a"] = len(rows)
    D.page_break()
    rows = []
    RSETS5 = ["U", "F", "K", "M", "M11s", "M19", "P10s", "N8s", "U+F", "U+K", "U+M", "U+M11s", "U+M19", "U+F+K+M",
              "U+F+K+M19"]
    for s in RSETS5:
        for r in WOUND:
            a, b = BE["rel"][s]["40"], BE["rel"][s]["11"]
            rows.append([bl(s), NICE[r], eiv(a[f"{r}|bench"]), eiv(a[f"{r}|diff"]), eiv(b[f"{r}|bench"]),
                         eiv(b[f"{r}|diff"]), eiv(a[f"{r}|other"])])
    D.caption("Table S5b", "Benchmark validity per wound reference",
              f"Generator B at texture amplitude 40 and 11, per wound reference R. {BENCHTXT} {BMB} The four "
              "references share data, so their values are not independent.")
    D.table(["Set", "Reference", "Bench, amplitude 40", "Diff, amplitude 40", "Bench, amplitude 11",
             "Diff, amplitude 11", "Other"], rows, [1.8, 2.6, 3.8, 3.8, 3.8, 3.8, 3.8], size=7.5, left_cols=2)
    counts["S5b"] = len(rows)
    D.page_break()
    rows = []
    for s in ("K", "M", "M11s", "M19", "P10s", "N8s"):
        src = BE["v6"] if s in ("K", "M") else BE
        for r in WOUND + ["mean"]:
            h4, h1 = src["rel_hier"][s]["40"], src["rel_hier"][s]["11"]
            d4, d1 = src["rel_dataonly"][s]["40"], src["rel_dataonly"][s]["11"]
            assert h4[f"{r}|bench"]["est"] == BE["rel"][s]["40"][f"{r}|bench"]["est"]
            rows.append([s, "mean of 4" if r == "mean" else NICE[r], eiv(h4[f"{r}|bench"]), eiv(h4[f"{r}|diff"]),
                         eiv(d4[f"{r}|diff"]), eiv(h1[f"{r}|bench"]), eiv(h1[f"{r}|diff"]), eiv(h4[f"{r}|other"])])
    D.caption("Table S5c", "Benchmark validity with hierarchical intervals (K and modern sets)",
              f"Same quantities as Tables S5a and S5b with hierarchical and data-only intervals. {BHB} {BENCHTXT}")
    D.table(["Set", "Reference", "Bench 40 (h)", "Diff 40 (h)", "Diff 40 (data-only)", "Bench 11 (h)",
             "Diff 11 (h)", "Other (h)"], rows, [1.4, 2.4] + [3.4] * 6, size=7.5, left_cols=2)
    counts["S5c"] = len(rows)
    D.page_break()
    se = BE["seed_effect_amp40"]
    rows = []
    for p in ("P10s", "P10_s0", "P10_s1", "P10_s2", "N8s", "N8_s0", "N8_s1", "N8_s2", "M19", "M19_s0", "M19_s1",
              "M19_s2", "M", "M11s"):
        v = se[p]
        rows.append([p, str(v["n_methods"]), eiv(v["bench"]), eiv(v["other"]), eiv(v["diff"])] +
                    [f(v["bench_by_ref"][r]) for r in WOUND])
    D.caption("Table S5d", "Benchmark validity in single-seed method sets (amplitude 40)",
              "Single-seed sets: P10_sk, the original architectures at training seed k (models that exist); N8_sk, the "
              "new architectures at seed k; M19_sk = P10_sk + N8_sk + nnU-Net; M = P10_s0 + nnU-Net. Seed-averaged sets "
              f"for comparison. Bench, other, diff: means over the {NW} wound references, estimate [interval]; the last "
              f"{NW} columns give bench per reference. {BENCHTXT.split(' ALL57')[0]} {BMB}")
    D.table(["Set", "n", "Bench", "Other", "Diff"] + [f"Bench, {NICE[r]}" for r in WOUND], rows,
            [1.8, 0.8, 3.5, 3.5, 3.5] + [2.1] * 4, size=7.5)
    counts["S5d"] = len(rows)
    D.page_break()
    rows = []
    for s in BSETS:
        t = BE["texture"][s]
        h = BE["texture_hier"].get(s) or (BE["v6"]["texture_hier"].get(s) if s in ("M", "K") else None)
        rows.append([bl(s)] + [eiv(t[r]) for r in WOUND + ["mean"]] + [eiv(h["mean"]) if h else "n/a"])
    D.caption("Table S5e", "Texture effect (amplitude 40 minus amplitude 11)",
              "Texture effect: benchmark agreement at texture amplitude 40 minus that at amplitude 11 (each averaged "
              "over the 3 synthetic seeds), per wound reference and as the mean over the 4 references, estimate "
              f"[interval]. {BMB} Last column: the mean with the hierarchical interval where per-image synthetic "
              f"scores exist (n/a otherwise). {BHB.split(' Only')[0]}")
    D.table(["Set"] + [NICE[r] for r in WOUND] + ["Mean of 4", "Mean of 4 (h)"], rows, [1.8] + [3.7] * 6, size=7.5)
    counts["S5e"] = len(rows)
    D.page_break()
    rows = []
    for s in BSETS:
        mm = BE["matched"][s]
        rows.append([bl(s)] + [f(mm[f"{g}_seedmean|mean4"]["est"]) for g in ("nrA", "nrB0", "nrC0", "nrC34")] +
                    [eiv(mm["B0-C34_seedmean|mean4"])])
    D.caption("Table S5f", "Generator families",
              "Agreement of each generator family with the clinical wound data: Spearman agreement of mean Dice on the "
              "synthetic set with each of the 4 wound datasets, averaged over the 4 datasets and over the 3 synthetic "
              "seeds. nrA: generator A, uniform drape; nrB0: generator B, texture ablated; nrC0: generator C, no noise; "
              f"nrC34: generator C, noise 3.4. Last column: nrB0 minus nrC34, estimate [interval]. {BMB}")
    D.table(["Set", "Generator A (nrA)", "Generator B, texture ablated (nrB0)", "Generator C, no noise (nrC0)",
             "Generator C, noise 3.4 (nrC34)", "nrB0 minus nrC34"], rows, [1.8, 3.4, 3.4, 3.4, 3.4, 4.2], size=7.5)
    counts["S5f"] = len(rows)
    rows = []
    for n, v in BE["lomo"]["M19"].items():
        rows.append(["none (all 19)" if n == "_all" else short(n), f(v["bench40_mean"]), f(v["pairs6"]),
                     f(v["rel40_mean"])])
    assert abs(BE["lomo"]["M19"]["_all"]["pairs6"] - ST["point"]["dice|M19"]["wound pairs"]) < 1e-12
    D.caption("Table S5g", "Leave one method out of M19: benchmark validity (amplitude 40)",
              f"Bench, wound pairs (other) and diff as means over the {NW} wound references, after dropping one method "
              f"from M19 ({SETN['M19'] - 1} remaining). Point estimates. Method keys as in Table S3b.")
    D.table(["Dropped method", "Bench", f"Wound pairs ({comb(NW, 2)})", "Diff"], rows, [5.0, 3.0, 3.0, 3.0],
            size=7.5, bold_rows=(0,))
    counts["S5g"] = len(rows)
    return counts

# ====================================================================== S6
RG = ["M19", "K11"]
B_R = int(re.search(r"B (\d+)", RO["definitions"]["method_bootstrap"]).group(1))
assert re.search(r"default_rng\(0\)", RO["definitions"]["method_bootstrap"])
K_MISSING = re.search(r"\[(.*)\]", RO["definitions"]["K11_vs_paper_K14"]).group(1).replace("'", "")
ROBTXT = (f"Modern sets from the rescore files of the trained models: M19 as in the main analysis; K11, the K set "
          f"without its {len(K_MISSING.split(','))} pixel classifiers ({K_MISSING}), which were not rescored. Post-processing "
          "P0: Gaussian smoothing of the probability map, threshold, small-component removal (main analysis); P1: "
          "threshold, closing and opening with a 5 px ellipse, small-component removal; P2: threshold and "
          "small-component removal only. nnU-Net values in the rescore files come from the second-session "
          f"re-prediction (Table S9). Intervals: method bootstrap, B = {B_R:,}, 95% percentile, the same draws for every "
          "condition, so differences between conditions get paired intervals.")
SUMQ = [("wound pairs", "Wound pairs"), ("colonoscopy pairs", "Colonoscopy pairs"),
        ("modality_diff", "Modality difference (55)")]


def table_s6(D):
    counts = {}
    rows = []
    for g in RG:
        for m, mn in (("dice", "Dice"), ("nsd2", "NSD")):
            for v in ("P0", "P1", "P2"):
                e = RO["postprocessing"]["summaries"][f"{g}|{m}|{v}"]
                dl = "" if v == "P0" else eiv(RO["postprocessing"]["delta_vs_P0"][f"{g}|{m}|{v}-P0|wound pairs"])
                rows.append([g, mn, v] + [eiv(e[k]) for k, _ in SUMQ] + [dl])
    D.caption("Table S6a", "Robustness to post-processing (modern sets)",
              f"{ROBTXT} Last column: change in wound-pair agreement from P0 with its paired interval. NSD at 2 px. "
              f"{PAIRTXT.split(' Control')[0]}", bm="S6", before=0)
    D.table(["Set", "Metric", "Variant"] + [h for _, h in SUMQ] + ["Wound pairs minus P0"], rows,
            [1.3, 1.3, 1.3, 4.3, 4.3, 4.3, 4.3], size=7.5, left_cols=3)
    counts["S6a"] = len(rows)
    rows = []
    TM = [("dice", "Dice"), ("nsd1", "NSD 1 px"), ("nsd2", "NSD 2 px"), ("nsd3", "NSD 3 px"), ("nsd5", "NSD 5 px"),
          ("nsd_diag1", "NSD 1% of diagonal")]
    for g in RG:
        for m, mn in TM:
            e = RO["tolerance"][f"{g}|{m}|P0"]
            dl = "" if m == "nsd2" else eiv(RO["tolerance_delta_vs_nsd2"][f"{g}|{m}-nsd2|wound pairs"])
            rows.append([g, mn] + [eiv(e[k]) for k, _ in SUMQ] + [dl])
    D.caption("Table S6b", "Robustness to the NSD tolerance (modern sets, P0)",
              "NSD at a boundary tolerance of 1, 2 (main analysis), 3 and 5 px and at 1% of the image diagonal; Dice "
              "for reference. Last column: change in wound-pair agreement from NSD at 2 px, paired interval. "
              f"{ROBTXT.split(' Post-processing')[0]} Intervals: method bootstrap, B = {B_R:,}, 95% percentile.")
    D.table(["Set", "Metric"] + [h for _, h in SUMQ] + ["Wound pairs minus NSD 2 px"], rows,
            [1.3, 2.8, 4.3, 4.3, 4.3, 4.3], size=7.5, left_cols=2)
    counts["S6b"] = len(rows)
    D.page_break()
    TH = RO["threshold"]
    grid = sorted({k.split("|")[1] for k in TH["common"]}, key=float)
    rows = []
    for g in RG:
        e = TH["default"][g]
        rows.append([g, "default (own FUSeg-tuned threshold)"] + [eiv(e[k]) for k, _ in SUMQ] + [""])
        e = TH["oracle"][g]
        rows.append([g, "oracle per dataset"] + [eiv(e[k]) for k, _ in SUMQ] +
                    [eiv(TH["delta_vs_default"][f"{g}|oracle|wound pairs"])])
    D.caption("Table S6c", "Threshold protocol: default and oracle (Dice, P0)",
              "Default: each model at its own threshold tuned on held-out FUSeg training images (main analysis). "
              "Oracle: for each individual model and dataset, the grid threshold with the best mean Dice on that "
              "dataset (chosen on the scored images, an optimistic upper bound); seed averages average their members' "
              f"oracle scores. Grid: {f(float(grid[0]), 2)} to {f(float(grid[-1]), 2)} in steps of "
              f"{f(float(grid[1]) - float(grid[0]), 2)}. Last column: change in wound-pair agreement from the default, "
              f"paired interval. {ROBTXT.split(' Post-processing')[0]} Intervals: method bootstrap, B = {B_R:,}, 95% percentile.")
    D.table(["Set", "Threshold"] + [h for _, h in SUMQ] + ["Wound pairs minus default"], rows,
            [1.3, 4.6, 4.0, 4.0, 4.0, 4.0], size=7.5, left_cols=2)
    counts["S6c"] = len(rows)
    rows = []
    pd_ = TH["calibration_probe"]["M19|default"]
    rows.append(["default", eiv(TH["default"]["M19"]["wound pairs"]), "", eiv(TH["default"]["M19"]["modality_diff"]),
                 f(float(np.mean(list(pd_["group_mean_dice"].values())))), f(pd_["mean_rho_wound"]),
                 eiv(TH["default"]["K11"]["wound pairs"]), eiv(TH["default"]["K11"]["modality_diff"])])
    for t in grid:
        e, k_ = TH["common"][f"M19|{t}"], TH["common"][f"K11|{t}"]
        pr = TH["calibration_probe"][f"M19|{t}"]
        rows.append([t, eiv(e["wound pairs"]), eiv(TH["delta_vs_default"][f"M19|{t}|wound pairs"]),
                     eiv(e["modality_diff"]), f(float(np.mean(list(pr["group_mean_dice"].values())))),
                     f(pr["mean_rho_wound"]), eiv(k_["wound pairs"]), eiv(k_["modality_diff"])])
    assert set(pd_["group_mean_dice"]) == set(WOUND)
    D.caption("Table S6d", "One common threshold for all models (Dice, P0)",
              "Every model thresholded at the same grid value. Mean wound Dice: M19 group mean Dice averaged over the "
              f"{NW} wound datasets. Calibration rho: mean over the {NW} wound datasets of the Spearman correlation "
              "between the M19 method means and each method's default (FUSeg-tuned) threshold; a high value means that "
              "the ranking follows how confident each model's probabilities are. Minus default: change in M19 "
              f"wound-pair agreement from the default protocol, paired interval. Intervals: method bootstrap, "
              f"B = {B_R:,}, 95% percentile.")
    D.table(["Threshold", "M19 wound pairs", "Minus default", "M19 modality difference", "Mean wound Dice",
             "Calibration rho", "K11 wound pairs", "K11 modality difference"], rows,
            [1.7, 3.6, 3.6, 3.6, 1.8, 1.8, 3.6, 3.6], size=7.5)
    counts["S6d"] = len(rows)
    D.page_break()
    rows = []
    for d in DS:
        o = TH["oracle_detail"]["M19"][d]
        rows.append([SHORTNICE[d], f(o["group_mean_default"]), f(o["group_mean_oracle"]),
                     pc(o["share_models_oracle_thr_differs"]), pc(o["share_models_oracle_at_grid_min_0.05"]),
                     f(o["median_oracle_thr"], 2), f(o["spearman_default_vs_oracle_methods"])])
    assert grid[0] == "0.05"
    D.caption("Table S6e", "Oracle thresholds per dataset (M19)",
              "Group mean Dice at the default and at the oracle threshold; share of the individual M19 models whose "
              "oracle threshold differs from their default; share at the lowest grid value (0.05, where the oracle is "
              "limited by the grid); median oracle threshold; Spearman correlation of the M19 method means at default "
              "and oracle thresholds.")
    D.table(["Dataset", "Mean Dice, default", "Mean Dice, oracle", "Share with other threshold",
             "Share at grid minimum", "Median oracle threshold", "Rank correlation default with oracle"], rows,
            [3.2, 2.6, 2.6, 2.8, 2.8, 2.8, 3.4], size=7.5)
    counts["S6e"] = len(rows)
    rows = []
    for g in RG:
        for m, mn in (("dice", "Dice"), ("nsd2", "NSD")):
            e = RO["resolution"][f"{g}|{m}"]
            rows.append([g, mn, f(e["fuseg_mean"]), f(e["fuseg_r256_mean"]),
                         f"{fs(e['shift_mean'])} [{fs(e['shift_min'])}, {fs(e['shift_max'])}]",
                         f(e["spearman_fuseg_vs_r256"]), eiv(e["with_fuseg"]["wound pairs"]),
                         eiv(e["with_fuseg_r256"]["wound pairs"]), eiv(e["delta_r256_minus_native_wound_pairs"])])
    D.caption("Table S6f", "Resolution: FUSeg at native size and at 256 px (P0)",
              "FUSeg at 256 px: the same test images and masks resampled so that the long side is 256 px (image by area "
              "averaging, mask by nearest neighbour), with the FUSeg exclusions. Shift: 256 px minus native, mean over "
              "the methods of the set [smallest, largest]. Rank correlation: Spearman correlation of the method means at "
              "the two resolutions. Wound pairs computed with FUSeg at native size or at 256 px, and the change with its "
              f"paired interval. Intervals: method bootstrap, B = {B_R:,}, 95% percentile.")
    D.table(["Set", "Metric", "Mean, native", "Mean, 256 px", "Shift [min, max]", "Rank correlation",
             "Wound pairs, native", "Wound pairs, 256 px", "Change"], rows,
            [1.2, 1.2, 1.8, 1.8, 3.6, 1.9, 3.6, 3.6, 3.6], size=7.5, left_cols=2)
    counts["S6f"] = len(rows)
    if RP is not None:
        counts.update(table_s6_pooled(D))
    return counts


CONDLAB = {"dice|P0": "Dice, P0 (main)", "dice|P1": "Dice, P1", "dice|P2": "Dice, P2", "nsd2|P0": "NSD 2 px, P0 (main)",
           "nsd2|P1": "NSD 2 px, P1", "nsd2|P2": "NSD 2 px, P2", "nsd1|P0": "NSD 1 px", "nsd3|P0": "NSD 3 px",
           "nsd5|P0": "NSD 5 px", "nsd_diag1|P0": "NSD 1% of diagonal", "dice|P0|fuseg_r256": "Dice, FUSeg at 256 px",
           "nsd2|P0|fuseg_r256": "NSD 2 px, FUSeg at 256 px"}


def yn(b):
    return "yes" if b else "no"


def table_s6_pooled(D):
    counts = {}
    RR = RP["results"]
    conds = list(RP["definitions"]["conditions"])
    assert conds == list(CONDLAB), conds
    PS = ["U+M19", "U+F", "U+K11"]
    assert RP["definitions"]["sets"]["U+K11"] == ["U", "K11"]
    B_RP = int(re.search(r"B (\d+)", RP["definitions"]["method_bootstrap"]).group(1))
    # cross-checks: modern rows equal robust_v7; main-condition pooled values equal stats_v7 (run-1 scores)
    for g in ("M19", "K11"):
        for m in ("dice", "nsd2"):
            for v in ("P0", "P1", "P2"):
                a = RR[g][f"{m}|{v}"]["pooled"]["wound pairs"]
                b = RO["postprocessing"]["summaries"][f"{g}|{m}|{v}"]["wound pairs"]
                assert a["est"] == b["est"] and a["ci"] == b["ci"], (g, m, v)
    for s, st in (("U+M19", "U+M19"), ("U+F", "U+F"), ("U", "U"), ("F", "F")):
        for m, sm in (("dice", "dice"), ("nsd2", "nsd")):
            a = RR[s][f"{m}|P0"]["pooled"]["wound pairs"]["est"]
            assert abs(a - ST["point"][f"{sm}|{st}"]["wound pairs"]) < 1e-12, (s, m)
    for s in PS:
        assert RR[s]["dice|P0"]["n_methods"] == sum(RP["definitions"]["groups"][g]["n"] for g in RP["definitions"]["sets"][s])
    ntxt = ", ".join(f"{s} ({RR[s]['dice|P0']['n_methods']})" for s in PS)
    base = (f"Pooled sets {ntxt}: U+M19 and U+F as in the main analysis; U+K11 uses K11, the K set without its "
            f"{len(K_MISSING.split(','))} pixel classifiers ({K_MISSING}), which were not rescored, so it differs from "
            "U+K. Scores from the rescore files of every method (untrained and from-scratch methods rescored locally, "
            "trained Kaggle models from the second session, nnU-Net re-predicted). Conditions: post-processing P0 "
            "(main), P1, P2 for Dice and NSD at 2 px; NSD at 1, 3, 5 px and at 1% of the image diagonal (P0); FUSeg "
            "replaced by FUSeg at 256 px (P0). For untrained methods P1 equals P0, because their main pipeline already "
            "applies closing and opening. Stratified agreement: methods ranked within their own group, ranks turned "
            "into within-group percentiles, Pearson correlation of the pooled percentiles. Size-weighted within: the "
            "size-weighted mean of the within-group agreements. Intervals: method bootstrap (methods resampled within "
            f"each group), B = {B_RP:,}, 95% percentile, the same draws for every condition (paired changes).")
    rows = []
    for s in PS:
        for c in conds:
            x = RR[s][c]["composition"]["wound"]
            other = [g for g in RP["definitions"]["sets"][s] if g != "U"][0]
            assert x["holds"] == (x["C1_pooled_above_both_within_ci_above0"] and x["C2_pooled_vs_other_ci_below0"]
                                  and x["C3_stratified_within_tol_of_size_weighted"])
            rows.append([s, CONDLAB[c], eiv(x["pooled"]), f(x["within"]["U"]), f"{other} {f(x['within'][other])}",
                         eiv(x["pooled_minus_higher_within"]), eiv(x["pooled_vs_other"]), eiv(x["stratified"]),
                         f(x["size_weighted_within"]),
                         f"{yn(x['C1_pooled_above_both_within_ci_above0'])} / {yn(x['C2_pooled_vs_other_ci_below0'])} / "
                         f"{yn(x['C3_stratified_within_tol_of_size_weighted'])}", yn(x["holds"])])
    tol = re.search(r"<= ([0-9.]+)", RP["definitions"]["composition_criteria"]["C3"]).group(1)
    D.page_break()
    D.caption("Table S6g", "Composition effect in pooled sets under every condition (wound domain)",
              f"{base} Gap: pooled wound-pair agreement minus the higher of the two within-group values (paired "
              f"interval). Criteria: C1, gap interval above 0; C2, interval of the pooled wound with non-wound agreement "
              f"below 0; C3, stratified agreement within {tol} of the size-weighted within-group value. Holds: all three. "
              f"{PAIRTXT.split(' Colonoscopy')[0]} Wound with non-wound: mean over the {NW * (len(REFS) - NW)} pairs of a "
              "wound dataset with a non-wound dataset.")
    D.table(["Set", "Condition", "Pooled wound pairs", "Within U", "Within other group", "Gap", "Pooled wound with "
             "non-wound", "Stratified wound pairs", "Size-weighted within", "C1 / C2 / C3", "Holds"], rows,
            [1.5, 3.0, 3.0, 1.3, 2.0, 3.0, 3.0, 3.0, 1.6, 2.2, 1.1], size=7, left_cols=2)
    counts["S6g"] = len(rows)
    rows = []
    for s in PS:
        for c in conds:
            x = RR[s][c]["composition"]["colonoscopy"]
            other = [g for g in RP["definitions"]["sets"][s] if g != "U"][0]
            rows.append([s, CONDLAB[c], eiv(x["pooled"]), f(x["within"]["U"]), f"{other} {f(x['within'][other])}",
                         eiv(x["pooled_minus_higher_within"]), eiv(x["pooled_vs_other"]), eiv(x["stratified"]),
                         f"{yn(x['C1_pooled_above_both_within_ci_above0'])} / {yn(x['C2_pooled_vs_other_ci_below0'])} / "
                         f"{yn(x['C3_stratified_within_tol_of_size_weighted'])}", yn(x["holds"])])
    D.page_break()
    D.caption("Table S6h", "Composition criteria in the colonoscopy domain (mirror view)",
              f"The criteria of Table S6g applied with the {comb(NC, 2)} colonoscopy pairs and the "
              f"{NC * (len(REFS) - NC)} pairs of a colonoscopy dataset with a non-colonoscopy dataset in place of the "
              "wound pairs; relevant mainly for U+K11, whose trained methods were trained on colonoscopy images. "
              + base.split(" Conditions:")[0] + " Intervals as in Table S6g.")
    D.table(["Set", "Condition", "Pooled colonoscopy pairs", "Within U", "Within other group", "Gap",
             "Pooled colonoscopy with other", "Stratified colonoscopy pairs", "C1 / C2 / C3", "Holds"], rows,
            [1.5, 3.0, 3.2, 1.3, 2.0, 3.2, 3.2, 3.2, 2.2, 1.1], size=7, left_cols=2)
    counts["S6h"] = len(rows)
    rows = []
    for s in ["U", "F", "M19", "K11"] + PS:
        for c in conds:
            x = RR[s][c]
            chk = [k for k in x if k.startswith("change_vs_")]
            ch = x[chk[0]] if chk else None
            e = x["exact_modality"]
            assert e["n_relabelings"] == N_RELAB
            rows.append([s, CONDLAB[c], eiv(x["pooled"]["wound pairs"]),
                         eiv(ch["wound pairs"]) if ch else "",
                         eiv(ch["stratified wound pairs"]) if ch and "stratified wound pairs" in ch else "",
                         eiv(x["pooled"]["modality_diff"]), fp(e["p"], e["min_p"])])
    ref = RP["definitions"]["change_reference"]
    grp = {}
    for k, v in ref.items():
        grp.setdefault(v, []).append(CONDLAB[k])
    reftxt = "; ".join(f"{CONDLAB[v]} for {', '.join(ks)}" for v, ks in grp.items())
    D.page_break()
    D.caption("Table S6i", "Wound-pair agreement, changes and modality difference under every condition",
              f"Single groups (U, F, M19, K11) and pooled sets. Changes are relative to the main condition of the same "
              f"metric, with paired intervals ({reftxt}). Stratified "
              f"change: pooled sets only. Modality difference over the 55 pairs with its interval and exact p. {EXTXT} "
              + base.split(" Stratified agreement:")[0] + f" Intervals: method bootstrap, B = {B_RP:,}, 95% percentile.")
    D.table(["Set", "Condition", "Wound pairs", "Change in wound pairs", "Change in stratified wound pairs",
             "Modality difference (55)", "Exact p"], rows, [1.5, 3.4, 3.6, 3.6, 3.6, 3.6, 2.2], size=7, left_cols=2)
    counts["S6i"] = len(rows)
    return counts

# ====================================================================== S7
SEL = PR["selection"]
MISS = PR["miss"]["0.1"]
SETS5 = WOUND + [CONTROL]
B_P = PR["B"]
UBTXT = (f"Intervals: unit bootstrap (resampling units drawn with replacement, the same draws for every model of a "
         f"dataset, models fixed), B = {B_P:,}, 95% percentile; they cover image sampling only. {UNITTXT}")
CANDTXT = (f"Candidate sets: M19 ({SEL['M19']['n_candidates']}, seed averages plus nnU-Net); M ({SEL['M']['n_candidates']}, "
           f"original architectures at seed 0 plus nnU-Net); M11s ({SEL['M11s']['n_candidates']}); P10s "
           f"({SEL['P10s']['n_candidates']}) and N8s ({SEL['N8s']['n_candidates']}), the original and new architectures "
           "alone as seed averages; M19_sk: every architecture that has training seed k, as that single model, plus "
           f"nnU-Net ({SEL['M19_s0']['n_candidates']}, {SEL['M19_s1']['n_candidates']}, {SEL['M19_s2']['n_candidates']} "
           "candidates).")
LOSSTXT = ("Selection loss from A to B: best mean Dice on B minus the mean Dice on B of the candidate with the highest "
           "mean Dice on A. Random choice: mean over candidates of (best on B minus candidate on B). Synthetic chooser: "
           "the candidate with the highest mean Dice on generator B at the given texture amplitude (mean over the 3 "
           "synthetic seeds). Point estimates (a loss depends on a maximum over candidates, so no interval is given).")


def table_s7(D):
    counts = {}
    s19 = SEL["M19"]
    rows = []
    for a in WOUND:
        L = s19["matrix_rows_choose_on_cols_deploy_on"][a]
        rows.append([f"{NICE[a]} ({short(s19['chosen_on'][a])})"] + [f(L[b]) for b in WOUND] +
                    [f(float(np.mean([L[b] for b in WOUND if b != a])))])
    rows.append([f"Control ({short(s19['control_chooser'])})"] + [f(s19["control_chooser_loss"][b]) for b in WOUND] +
                [f(float(np.mean([s19["control_chooser_loss"][b] for b in WOUND])))])
    for amp in ("amp11", "amp28", "amp40"):
        v = s19["synthetic"][amp]
        assert abs(np.mean([v["loss"][b] for b in WOUND]) - v["mean_loss"]) < 1e-12
        rows.append([f"Generator B, amplitude {amp[3:]} ({short(v['chosen'])})"] + [f(v["loss"][b]) for b in WOUND] +
                    [f(v["mean_loss"])])
    rows.append(["Random choice (expected)"] + [f(s19["random_choice_expected_loss"][b]) for b in WOUND] +
                [f(s19["random_mean_over_B"])])
    off = [s19["matrix_rows_choose_on_cols_deploy_on"][a][b] for a in WOUND for b in WOUND if a != b]
    assert abs(np.mean(off) - s19["offdiag_mean"]) < 1e-12
    D.caption("Table S7a", "Selection loss among the 19 modern methods (M19, Dice)",
              f"Rows: the set the model is chosen on (chosen model in brackets); columns: the wound dataset it is "
              f"deployed on. Mean: over the deployment datasets other than the choice dataset (for the wound rows) or "
              f"over all {NW} (other rows). Mean over the {len(off)} off-diagonal pairs: {f(s19['offdiag_mean'])} "
              f"(range {f(s19['offdiag_min'])} to {f(s19['offdiag_max'])}). {LOSSTXT} Model keys as in Table S3b.",
              bm="S7", before=0)
    D.table(["Chosen on (pick)"] + [f"Deployed on {NICE[b]}" for b in WOUND] + ["Mean"], rows,
            [8.0, 3.2, 3.2, 3.2, 3.2, 2.4], size=7.5)
    counts["S7a"] = len(rows)
    rows = []
    for k in ("M", "M11s", "M19", "M19_s0", "M19_s1", "M19_s2", "P10s", "N8s"):
        v = SEL[k]["synthetic"]
        a4 = v["amp40"]
        rows.append([k, str(SEL[k]["n_candidates"]), short(v["amp11"]["chosen"]), short(v["amp28"]["chosen"]),
                     short(a4["chosen"]), " / ".join(short(a4["chosen_per_synth_seed"][s]) for s in ("2024", "7", "101")),
                     " / ".join(f(a4["loss"][b]) for b in WOUND), f(a4["mean_loss"]),
                     f"{a4['n_within_0.01']} / {a4['n_within_0.02']}", f(SEL[k]["offdiag_mean"]),
                     f(SEL[k]["random_mean_over_B"])])
    D.caption("Table S7b", "Synthetic chooser by candidate set and training seed",
              f"{CANDTXT} Pick: the synthetic chooser's model at amplitude 11, 28 and 40, and at amplitude 40 per "
              "synthetic seed (2024 / 7 / 101). Loss per wound dataset at amplitude 40 in the order Medetec / "
              "CO2Wounds-V2 / FUSeg / AZH, and its mean. Within 0.01 / 0.02: number of the 4 wound datasets where the "
              "amplitude-40 pick is within 0.01 or 0.02 Dice of the best. Other wound dataset: mean loss when choosing "
              f"on one wound dataset and deploying on another (12 ordered pairs). {LOSSTXT}")
    D.table(["Candidate set", "n", "Pick, amplitude 11", "Pick, amplitude 28", "Pick, amplitude 40",
             "Pick per synthetic seed (amplitude 40)", "Loss per wound dataset (amplitude 40)", "Mean loss",
             "Within 0.01 / 0.02", "Other wound dataset", "Random choice"], rows,
            [1.6, 0.7, 2.3, 2.3, 2.3, 4.4, 4.0, 1.4, 1.5, 1.6, 1.5], size=7, left_cols=7)
    counts["S7b"] = len(rows)
    D.page_break()
    cb = PR["clinical_best_by_training_seed"]
    rows = [[SHORTNICE[s]] + [f"{short(cb[s][k]['model'])} ({f(cb[s][k]['mean_dice'])})"
                              for k in ("M19", "M19_s0", "M19_s1", "M19_s2")]
            + [f"{short(SEL['M']['best_per_dataset'][s]['model'])} ({f(SEL['M']['best_per_dataset'][s]['mean_dice'])})"]
            for s in SETS5]
    D.caption("Table S7c", "Best model per dataset by training seed",
              f"Model with the highest mean Dice on each wound dataset and the control (mean Dice in brackets), among "
              f"the seed-averaged M19 candidates, the single-seed candidate sets M19_s0, M19_s1, M19_s2, and M. {CANDTXT}")
    D.table(["Dataset", "M19 (seed averages)", "M19_s0", "M19_s1", "M19_s2", "M"], rows,
            [2.4, 4.4, 4.4, 4.4, 4.4, 4.4], size=7.5, left_cols=6)
    counts["S7c"] = len(rows)
    D.page_break()
    rows = []
    for a, z0 in MISS["azh"]["per_architecture"].items():
        z = {s: MISS[s]["per_architecture"][a] for s in SETS5}
        grp = "nnU-Net" if a == "nnunet_2d" else ("original" if a.endswith("_in") else "new")
        rows.append([short(a), grp, str(z["azh"]["n_seeds"]),
                     f"{pc(z['azh']['mean'])} ({pc(z['azh']['min'])} to {pc(z['azh']['max'])})\n{pciv(z['azh']['ci_of_seed_mean'])}",
                     f"{pc(z['fuseg']['mean'])} ({pc(z['fuseg']['min'])} to {pc(z['fuseg']['max'])})",
                     f"{pc(z[CONTROL]['mean'])} ({pc(z[CONTROL]['min'])} to {pc(z[CONTROL]['max'])})",
                     pc(z["medetec"]["mean"]), pc(z["co2wounds"]["mean"])])
    D.caption("Table S7d", "Misses per architecture",
              "Miss: an image with a non-empty reference mask and Dice below 0.1. Miss rate: misses over non-empty "
              f"images (AZH {MISS['azh']['n_nonempty']} images in {MISS['azh']['n_units_nonempty']} photographs; FUSeg "
              f"{MISS['fuseg']['n_nonempty']}; control {MISS[CONTROL]['n_nonempty']} in {MISS[CONTROL]['n_units_nonempty']} "
              f"photographs; Medetec {MISS['medetec']['n_nonempty']}; CO2Wounds-V2 {MISS['co2wounds']['n_nonempty']}). "
              "Computed for every individual model; an architecture's rate is the mean over its training seeds (smallest "
              "to largest single-seed rate in brackets). A seed-averaged Dice is not used for misses because it is not "
              f"the output of any model. AZH: second line gives the interval of the seed mean. {UBTXT}")
    D.table(["Architecture", "Set", "Seeds", "AZH: mean (min to max), interval", "FUSeg: mean (min to max)",
             "Control: mean (min to max)", "Medetec", "CO2Wounds-V2"], rows,
            [3.0, 1.6, 1.1, 5.0, 4.4, 4.4, 1.8, 2.2], size=7.5, left_cols=2)
    counts["S7d"] = len(rows)
    D.page_break()
    rows = []
    for s in ("azh", "fuseg", CONTROL):
        for gname, lab in (("P10_archmean", "original 10"), ("N8_archmean", "new 8"), ("P18_archmean", "all 18")):
            v = MISS[s]["groups"][gname]
            rows.append([lab, SHORTNICE[s], f"{pc(v['est'])} {pciv(v['ci'])}",
                         f"{pc(v['minus_nnunet'])} {pciv(v['minus_nnunet_ci'])}",
                         f"{pc(v['arch_mean_range'][0])} to {pc(v['arch_mean_range'][1])}",
                         f"{pc(v['single_model_range'][0])} to {pc(v['single_model_range'][1])}"])
        nn = MISS[s]["per_model"]["nnunet_2d"]
        rows.append(["nnU-Net", SHORTNICE[s], f"{pc(nn['miss_rate'])} {pciv(nn['ci'])}", "", "", ""])
    for s in ("azh", "fuseg", CONTROL):
        gg = MISS[s]["groups"]
        for k, lab in (("P10_s0_modelmean", "original, seed 0"), ("P10_s1_modelmean", "original, seed 1"),
                       ("P10_s2_modelmean", "original, seed 2"), ("N8_s0_modelmean", "new, seed 0"),
                       ("N8_s1_modelmean", "new, seed 1"), ("N8_s2_modelmean", "new, seed 2")):
            rows.append([lab, SHORTNICE[s], f"{pc(gg[k]['est'])} {pciv(gg[k]['ci'])}", "", "", ""])
    D.caption("Table S7e", "Misses by group and by training seed",
              "Group miss rate: mean over architectures, each architecture weighted equally (rows original 10, new 8, all "
              "18) or mean over the models of one training seed (rows by seed). Minus nnU-Net: paired difference to "
              f"nnU-Net from the same draws. Ranges: over architecture means and over single models. {UBTXT}")
    D.table(["Group", "Dataset", "Mean miss rate [interval]", "Minus nnU-Net [interval]", "Range of architecture means",
             "Range of single models"], rows, [3.4, 2.2, 4.4, 4.4, 4.0, 4.0], size=7.5, left_cols=2)
    counts["S7e"] = len(rows)
    D.page_break()
    pdg = PR["padding"]
    rows = []
    for s in ("azh", CONTROL, "fuseg"):
        for gname, lab in (("P10_all_seeds", "original 10, all seeds"), ("N8_all_seeds", "new 8, all seeds"),
                           ("P18_all_seeds", "all 18, all seeds")):
            v = pdg[s][gname]
            rows.append([SHORTNICE[s], f"{lab} ({v['n_models']})", f(v["mean_miss_share"]),
                         f(v["spearman_black_share_vs_miss_share"], 2), f(v["spearman_frac_vs_miss_share"], 2),
                         str(v["n_majority_miss_nnunet_detect"]), str(v["n_nnunet_miss_majority_detect"])])
    for s in ("azh",):
        for gname, lab in (("P18_all_seeds", "all 18"), ("N8_all_seeds", "new 8")):
            for t in pdg[s][gname]["miss_by_frac_tercile_and_padding"]:
                rows.append([SHORTNICE[s], f"{lab}, lesion-size tercile {t['frac_tercile'] + 1} (n {t['n']})",
                             f"{f(t['miss_more_padding'], 2)} vs {f(t['miss_less_padding'], 2)}", "", "", "", ""])
    D.caption("Table S7f", "Misses and image padding (all pretrained models)",
              "Miss share of an image: share of the individual pretrained models (number in brackets) that miss it. "
              "Black share: share of image pixels with every channel at most 8 (padding of the crop). Columns: mean miss "
              "share over non-empty images; Spearman correlation of the miss share with the black share and with the "
              "lesion fraction; number of images missed by most pretrained models but detected by nnU-Net, and the "
              "reverse. Last rows (AZH): mean miss share in images with more vs less padding than the median within each "
              f"lesion-size tercile (tercile 1 = smallest lesions). Mean black share over non-empty images: AZH "
              f"{f(pdg['azh']['black_share_mean_nonempty'])}, FUSeg {f(pdg['fuseg']['black_share_mean_nonempty'])}. "
              "Correlational: padding is not manipulated.")
    D.table(["Dataset", "Models", "Mean miss share", "Rank correlation with black share",
             "Rank correlation with lesion fraction", "Majority miss, nnU-Net detects", "nnU-Net misses, majority detects"],
            rows, [2.2, 6.0, 2.6, 3.0, 3.0, 3.0, 3.0], size=7.5, left_cols=2)
    counts["S7f"] = len(rows)
    return counts

# ====================================================================== S8
_RSRC = open(os.path.join(HERE, "robust_v7.py")).read()
B_U = int(re.search(r"^B = (\d+)$", _RSRC, re.M).group(1))
SEED_U = int(re.search(r"^SEED_UNIT = (\d+)$", _RSRC, re.M).group(1))
assert f"{B_U} draws, seed {SEED_U}" in open(os.path.join(V7, "ROBUST_V7.md")).read()
SAMLAB = [("SAM r1 (gt box)", "SAM, reference box"), ("MedSAM r1 (gt box)", "MedSAM, reference box"),
          ("MedSAM official (gt box)", "MedSAM official, reference box"), ("SAM (nnU-Net box)", "SAM, nnU-Net box"),
          ("MedSAM official (nnU-Net box)", "MedSAM official, nnU-Net box"), ("nnU-Net r1", "nnU-Net (first session)"),
          ("nnU-Net rescore P0", "nnU-Net (re-prediction)")]
SAMTXT = ("SAM and MedSAM (ViT-B) are prompted with boxes. Reference box: one box per connected component of the "
          "reference mask (an oracle prompt; these models are therefore never pooled with the other methods). MedSAM "
          "official: MedSAM with the preprocessing of its official inference script (second GPU session); MedSAM "
          "without that label is the first-session run. nnU-Net box: boxes of the nnU-Net prediction (up to 10 largest "
          "components); an image without an nnU-Net component gets no prompt and an empty prediction. Means over the "
          "kept images.")


def table_s8(D):
    counts = {}
    SM = RO["sam"]["means"]
    for part, m, mn in (("a", "dice", "Dice"), ("b", "nsd2", "NSD")):
        rows = [[SHORTNICE[d], str(SM[d]["n_images"])] + [f(SM[d][k][m]) for k, _ in SAMLAB] for d in DS]
        for d in DS:
            assert SM[d]["n_images"] == NKEPT[d]
            for k, n in (("SAM r1 (gt box)", "sam_vitb_box"), ("MedSAM r1 (gt box)", "medsam_vitb_box"),
                         ("MedSAM official (gt box)", "medsam_fixed_gtbox"), ("nnU-Net r1", "nnunet_2d")):
                assert abs(SM[d][k][m] - MEANS["dice" if m == "dice" else "nsd"][d][n]) < 1e-9, (d, k, m)
        D.caption(f"Table S8{part}", f"SAM and MedSAM, mean {mn}",
                  f"{SAMTXT} nnU-Net for comparison, from the first session and from the second-session re-prediction "
                  f"whose masks gave the nnU-Net boxes.{' NSD at 2 px.' if m == 'nsd2' else ''}",
                  bm="S8" if part == "a" else None, before=0 if part == "a" else 10)
        D.table(["Dataset", "n"] + [h for _, h in SAMLAB], rows, [3.0, 0.9] + [3.0] * 7, size=7.5)
        counts[f"S8{part}"] = len(rows)
    D.page_break()
    CMP = [("MedSAM official (gt box) - MedSAM r1 (gt box)", "MedSAM official minus MedSAM, reference box"),
           ("SAM (nnU-Net box) - nnU-Net rescore P0", "SAM with nnU-Net box minus nnU-Net"),
           ("MedSAM official (nnU-Net box) - nnU-Net rescore P0", "MedSAM official with nnU-Net box minus nnU-Net")]
    rows = []
    for d in DS:
        x = RO["sam"]["diffs"][d]
        rows.append([SHORTNICE[d]] + [eiv(x[f"{k}|{m}"]) for m in ("dice", "nsd2") for k, _ in CMP])
    D.caption("Table S8c", "Paired differences: official preprocessing and nnU-Net-box prompting",
              "Paired differences in mean Dice and mean NSD (2 px), estimate [interval]. nnU-Net-box variants are "
              "compared with the nnU-Net re-prediction whose masks gave the boxes. Intervals: unit bootstrap, "
              f"B = {B_U:,}, 95% percentile. {UNITTXT} {SAMTXT}")
    D.table(["Dataset"] + [f"Dice: {h}" for _, h in CMP] + [f"NSD: {h}" for _, h in CMP], rows,
            [2.8] + [3.8] * 6, size=7.5)
    counts["S8c"] = len(rows)
    rows = []
    for d in DS:
        z = RO["sam"]["no_box"][d]
        rows.append([SHORTNICE[d], str(z["n_images"]), str(z["empty_refs"]), str(z["empty_ref_and_nnunet_empty_exact"]),
                     str(z["nonempty_ref_dice0_allnsd0"]), str(z["of_which_nnbox_gt0_so_box_present"]),
                     f"{z['no_box_lower']} to {z['no_box_upper']}"])
    D.caption("Table S8d", "Images without an nnU-Net box",
              "An image without an nnU-Net component gets no prompt and an empty prediction, scored 1 against an empty "
              "reference and 0 otherwise. The nnU-Net masks were not stored, so the count is exact only for empty "
              "references. For non-empty references, nnU-Net Dice 0 with NSD 0 at every tolerance is necessary but not "
              "sufficient for an empty mask; candidates where an nnU-Net-box variant scores above 0 had a box and are "
              "removed, which gives the range in the last column.")
    D.table(["Dataset", "n", "Empty references", "No box, empty reference (exact)", "Non-empty candidates",
             "Candidates with a box", "No box (range)"], rows, [3.2, 1.2, 2.4, 3.4, 3.0, 3.0, 2.8], size=7.5)
    counts["S8d"] = len(rows)
    return counts


# ====================================================================== S9


def table_s9(D):
    counts = {}
    rows = []
    for d in DS:
        n_file = len(jl(RES, "perimage32", d + ".json")["ids"])
        ex = len(EXCL.get(d, []))
        assert n_file - ex == NKEPT[d] == ST["n_images"][d]
        rows.append([SHORTNICE[d], P7.MODALITY.get(d, "wound"), P7.SOURCE.get(d, "AZH"), str(n_file), str(ex),
                     str(NKEPT[d]), str(ST["n_units"][d]), str(ST["n_empty"][d])])
    dup = [e for e in EXCL["fuseg"] if e[2].startswith("byte-identical")]
    dup_tr = [e for e in dup if "FUSeg train" in e[2]]
    dup_ev = [e for e in dup if "evaluation image" in e[2]]
    crop = [e for e in EXCL[CONTROL] if e[2].startswith("crop of FUSeg")]
    cdup = [e for e in EXCL[CONTROL] if e[2].startswith("pixel-identical to AZH test")]
    acrop = [e for e in EXCL["azh"] if e[2].startswith("crop of FUSeg train")]
    assert len(dup) == len(dup_tr) + len(dup_ev) == len(EXCL["fuseg"])
    assert len(crop) + len(cdup) == len(EXCL[CONTROL]) and len(acrop) == len(EXCL["azh"])
    D.caption("Table S9a", "Images, exclusions and resampling units",
              f"Images in the scored file, images excluded, images kept for every analysis, resampling units and kept "
              f"images with an empty reference mask. Exclusions: {len(dup)} FUSeg validation images are byte-identical "
              f"to FUSeg training images ({len(dup_tr)}) or to another validation image ({len(dup_ev)}); {len(acrop)} AZH "
              f"test images are crops of FUSeg training photographs used to train the models; {len(crop)} control images "
              f"are crops of FUSeg photographs used in this study and {len(cdup)} are pixel-identical to AZH test images "
              f"(Table S9e). "
              f"{UNITTXT} Synthetic sets: {NKEPT[SYNTH[0]]} images each, no exclusions.", bm="S9", before=0)
    D.table(["Dataset", "Modality", "Source", "Images", "Excluded", "Kept", "Units", "Empty references"], rows,
            [4.4, 2.6, 2.4, 1.8, 1.8, 1.8, 1.8, 2.4], size=7.5, left_cols=3)
    counts["S9a"] = len(rows)
    nd = AU["check2_rescore_vs_scores"]["nnunet_detail"]
    rows = []
    for d in DS:
        z = nd["per_set"][d]
        rows.append([SHORTNICE[d], str(z["n"]), str(z["n_dice_gt_1e-4"]), str(z["n_dice_gt_1e-3"]),
                     f"{z['max_abs_dice']:.4f}", f"{z['max_abs_nsd2']:.4f}", f"{z['mean_dice_run1_kept']:.5f}",
                     f"{z['mean_dice_run2_kept']:.5f}", f"{z['delta_mean_dice_kept']:.1e}".replace("-", MINUS)])
    a = nd["all"]
    rows.append(["All", str(a["pairs"]), str(a["n_dice_gt_1e-4"]), str(a["n_dice_gt_1e-3"]), f"{a['max_abs_dice']:.4f}",
                 f"{a['max_abs_nsd2']:.4f}", "", "", ""])
    assert nd["rank_changes"] == []
    for d in DS:
        rr = nd["rank_among_11_v6_modern"][d]
        assert rr["rank_dice_run1"] == rr["rank_dice_run2"] and rr["rank_nsd2_run1"] == rr["rank_nsd2_run2"], d
    assert a["n_run2_higher"] + a["n_run2_lower"] == a["n_dice_gt_1e-4"]
    D.caption("Table S9b", "nnU-Net: first-session scores against the second-session re-prediction",
              "The same nnU-Net checkpoint predicted again in a second GPU session (half-precision inference). Images: "
              "per-image results compared (all images before exclusions). Counts of images whose Dice differs by more "
              "than 0.0001 and 0.001; largest absolute per-image difference in Dice and NSD (2 px); mean Dice over the "
              f"kept images in each session and the difference (second minus first). Direction of the "
              f"{a['n_dice_gt_1e-4']} larger differences: {a['n_run2_higher']} higher, {a['n_run2_lower']} lower in the "
              "second session. nnU-Net's rank among the 11 modern methods of M is the same in both sessions on every "
              "dataset, for Dice and NSD. The main analyses use the first-session scores.")
    D.table(["Dataset", "Images", "Dice differs > 0.0001", "Dice differs > 0.001", "Largest Dice difference",
             "Largest NSD difference", "Mean Dice, first session", "Mean Dice, second session", "Difference of means"],
            rows, [3.4, 1.4, 2.4, 2.4, 2.4, 2.4, 2.6, 2.6, 2.4], size=7.5, bold_rows=(len(rows) - 1,))
    counts["S9b"] = len(rows)
    D.page_break()
    ml = AU["check4_models"]["manet_log"]
    rows = []
    for n, v in ml.items():
        ep = {e["epoch"]: e for e in v["epochs_logged"]}
        assert sorted(ep) == [10, 20, 30, 40]
        nt = re.search(r"\((\d+) tensors", v["error"]) if v["error"] else None
        last = ep[40]["mean_loss"]
        rows.append([n] + [str(ep[k]["nonfinite_so_far"]) for k in (10, 20, 30, 40)] +
                     [f(last) if last is not None else "not finite",
                      f"not saved ({nt.group(1)} non-finite tensors)" if nt else "saved"])
    run2 = [r for r in AU["check4_models"]["rows"] if r.get("source") == "run 2" or "diverged" in r["status"]]
    saved = [r for r in run2 if r["status"] == "ok"]
    other_nf = [r["name"] for r in saved if r.get("nonfinite_steps")]
    assert other_nf == []
    D.caption("Table S9c", "Trainings that diverged",
              f"Second-session trainings: {len(run2)}; saved: {len(saved)}; every saved model had 0 non-finite training "
              "steps. MA-Net trainings with the cumulative number of optimiser steps with a non-finite loss at epochs 10, "
              "20, 30 and 40, the logged mean loss at epoch 40, and the outcome. A model with non-finite weights after "
              "training was not saved and is not imputed: manet_r18 is averaged over seeds 0 and 1, manet_r50 is a single "
              "model (seed 1).")
    D.table(["Model", "Non-finite steps, epoch 10", "Epoch 20", "Epoch 30", "Epoch 40", "Mean loss, epoch 40",
             "Outcome"], rows, [3.6, 2.6, 1.8, 1.8, 1.8, 2.6, 5.4], size=7.5)
    counts["S9c"] = len(rows)
    hit = [h for h in AU["check9_log"]["hits"]["FAILED"] if "woundseg" in h and "masks differ" in h]
    assert len(hit) == 1
    mm = re.search(r"(\d+) of (\d+) masks differ in size from their image \(> (\d+)%\)", hit[0])
    rows = [["WoundSeg", mm.group(2), mm.group(1), f"more than {mm.group(3)}%", "excluded before any scoring"]]
    D.caption("Table S9d", "Wound dataset excluded before scoring",
              "A fifth wound dataset was prepared and checked automatically before any model was run on it. Masks whose "
              "size differed from their image by more than the stated share indicate an unreliable image-mask pairing, "
              "so the dataset was excluded as a whole; no score exists for it.")
    D.table(["Dataset", "Image-mask pairs checked", "Masks not matching their image size", "Size difference",
             "Decision"], rows, [3.0, 3.6, 4.6, 3.0, 5.0], size=7.5)
    counts["S9d"] = len(rows)
    return counts

# ====================================================================== assemble Online Resource 1
TITLE = ("Method pool, clinical reference and training seed shape the measured validity of synthetic segmentation "
         "benchmarks: a study in wound imaging")
assert TITLE in open(os.path.join(MS, "BRIEF_MS_V6.md")).read()


def figure_s1(D):
    D.figure(FIGS1_PNG, 24.0)
    p = D.para(before=4, after=6)
    p.paragraph_format.line_spacing = 1.05
    D.run(p, "Figure S1. ", 9.5, bold=True)
    D.run(p, "Agreement matrices (Dice). ", 9.5, bold=True)
    D.run(p, (f"Each cell is the Spearman agreement of per-method mean Dice between the row dataset and the column "
              f"dataset over the methods of one set: (a) U, {SETN['U']} untrained methods; (b) M19, {SETN['M19']} modern "
              f"methods (seed averages and nnU-Net); (c) U+M19, the {SETN['U+M19']} methods pooled. Values to 2 decimals "
              "with the leading zero omitted (for example .54 is 0.54); negative values in bold; shading by absolute value "
              "(darker is stronger); diagonal not shown. Lines separate the wound, colonoscopy, dermoscopy and ultrasound "
              "datasets. Unrounded values for these and all other method sets are in Online Resource 2 (sheet Pair "
              "agreement)."), 9.5, color=GREYC)
    bm = "FigureS1"
    D.toc.append(("Figure S1", "Agreement matrices (Dice)", bm))
    D.bookmark(p, bm)


def build_or1(path):
    D = Doc()
    D.para("Online Resource 1", size=16, bold=True, after=6)
    D.para(TITLE, size=12, italic=True, after=12)
    D.para("Contents", size=11, bold=True, after=4)
    anchor = D.para(after=2)
    D.para("Notation", size=11, bold=True, before=10, after=4)
    D.para(("Numbers are rounded to 3 decimals (p values to 2 significant digits, percentages to 1 decimal); negative "
            "numbers use the minus sign. Intervals are 95% percentile bootstrap intervals; each caption names the "
            "bootstrap and its number of draws B. Datasets: " + DSNAMES + "; Control: the AZH training slice. "
            + SETTXT + " " + AGRTXT + " Method keys (for example unet_r18, segformer_mitb2) are listed with their "
            "architecture in Table S1; Online Resource 2 gives every per-method mean and every pairwise agreement."),
           size=9.5, color=GREYC, after=4)
    counts = {}
    sections = [table_s1, table_s2, table_s3, table_s4, table_s5, table_s6, table_s7, table_s8, table_s9]
    for fn in sections:
        D.page_break()
        c = fn(D)
        counts.update(c if isinstance(c, dict) else {"S1": c})
    D.page_break()
    figure_s1(D)
    # contents, inserted after the anchor in document order
    prev = anchor
    for label, title, bm in D.toc:
        top = re.fullmatch(r"Table S\d|Figure S1|Table S\da", label) is not None
        q = anchor.insert_paragraph_before()
        q.paragraph_format.space_after = Pt(1)
        q.paragraph_format.left_indent = Cm(0 if top and not label.endswith(("b", "c", "d", "e", "f", "g", "h", "i", "j"))
                                            else 0.6)
        D.link(q, f"{label}. {title}", bm, size=9.5)
    D.save(path)
    return counts


OR1 = os.path.join(MS, "Online_Resource_1.docx")
COUNTS = build_or1(OR1)
# house rules on every string written to the document
alltxt = "\n".join(TEXT_LOG)
assert "—" not in alltxt and "–" not in alltxt, "dash found"
bad = re.findall(r"(?<![\w.])-\d", alltxt)
assert not bad, bad[:5]
assert "fill in" not in alltxt.lower() and "todo" not in alltxt.lower() and "placeholder" not in alltxt.lower()
print("Online Resource 1 written:", json.dumps(COUNTS))
# the workbook: no dashes either
from openpyxl import load_workbook                     # noqa: E402
_wb = load_workbook(OR2)
for _ws in _wb:
    for _row in _ws.iter_rows(values_only=True):
        for _v in _row:
            if isinstance(_v, str):
                assert "\u2014" not in _v and "\u2013" not in _v, _v
print("robust_pooled_v7.json present:", RP is not None)
